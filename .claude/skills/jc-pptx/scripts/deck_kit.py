# -*- coding: utf-8 -*-
"""jc-pptx deck_kit v2 — 리멤버 웜 페이퍼 기본 테마 + 그리드·네비·컴포넌트 헬퍼 (python-pptx).

원칙: 좌표·색을 직접 하드코딩하지 말고 본 모듈의 상수·Theme·헬퍼를 쓴다.
지오메트리 정본: references/slide-types.md · 테마 매핑: references/themes.md
토큰 정본: jc-design-system/references/signature-tokens.md §6 (런타임 로드)

사용:
    from deck_kit import Deck, get_theme
    d = Deck(get_theme("remember"), doc_name="JLL CAN 2026 제안서", footer_text="JLL Client Appreciation Night 2026")
    d.cover("행사 기획·운영 제안서", "JLL Korea\\nClient Appreciation Night 2026", slogan="감사가 아니라, 내년의 첫 미팅", date="2026. 09. 23")
    s = d.content_slide("01 · 행사의 의미", "행사 이해", [("감사가 아니라, ", ""), ("'내년의 첫 미팅'", "accent")], sub="— 250명, 한 테이블")
    d.takeaway(s, [("좌석은 ", ""), ("의사결정 그룹", "deep"), ("에게만 씁니다.", "")])
    d.closing(contact=["mice_solution@remember.co.kr"])
    d.save("out.pptx")
"""
from __future__ import annotations

import glob
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

try:
    from PIL import Image
except Exception:  # Pillow 없으면 cover-crop·로고 비율 계산만 비활성
    Image = None

# ── 그리드 상수 (inch, 13.33 × 7.5) ─────────────────────────────
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.62
CONTENT_W = SLIDE_W - 2 * MARGIN          # 12.09
EYEBROW_Y, HEADLINE_Y, SUB_Y = 0.50, 0.90, 2.05
CONTENT_Y, CONTENT_BOTTOM = 2.50, 6.30
TAKEAWAY_Y, FOOTER_Y = 6.34, 7.06
GAP = 0.33

# 타이포 스케일 (pt) — jc-design-system §6 deckPt
SZ_COVER, SZ_KPI, SZ_HEADLINE, SZ_TITLE = 44, 48, 32, 22
SZ_SUB, SZ_BODY, SZ_SMALL, SZ_EYEBROW, SZ_CAPTION = 15, 13, 11, 11, 9
SZ_CARD_TITLE, SZ_LABEL, SZ_FOOTER = 15, 9, 9

HERE = Path(__file__).resolve()
ASSETS = HERE.parents[1] / "assets"


# ── SoT 탐색 ─────────────────────────────────────────────────────
def find_sot() -> Path | None:
    cands = [HERE.parents[2] / "jc-design-system",
             Path.home() / ".claude" / "skills" / "jc-design-system"]
    cands += sorted(Path.home().glob(".claude/skills/synced/*/jc-design-system"))
    for c in cands:
        if (c / "references" / "signature-tokens.md").is_file():
            return c
    return None


def _load_jc_tokens_module(sot: Path):
    import importlib.util
    p = sot / "scripts" / "jc_tokens.py"
    spec = importlib.util.spec_from_file_location("jc_tokens", p)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


# ── 테마 ───────────────────────────────────────────────────────
@dataclass
class Theme:
    name: str
    bg_base: str; bg_card: str
    text_primary: str; text_body: str; text_muted: str
    accent: str; accent_sub: str; line: str
    font_head: str = "Pretendard"
    font_body: str = "Pretendard"
    dark: bool = False
    texture: str = ""
    # v2 확장 (remember)
    surface_alt: str = ""; text_sub: str = ""
    accent_deep: str = ""; accent_light: str = ""; accent_pale: str = ""; accent_tint: str = ""
    line_soft: str = ""; charcoal: str = ""
    d_bg: str = ""; d_panel: str = ""; d_surface: str = ""; d_border: str = ""
    d_ink: str = ""; d_muted: str = ""; d_sub: str = ""; d_dim: str = ""; d_accent_text: str = ""
    positive: str = ""; negative: str = ""
    series: list = field(default_factory=list)
    logo_light: str = ""; logo_dark: str = ""; objet_dir: str = ""

    def __post_init__(self):
        # 확장 필드 미지정 시 기본 8토큰에서 유도 (명명 프리셋 호환)
        self.surface_alt = self.surface_alt or self.bg_card
        self.text_sub = self.text_sub or self.text_muted
        self.accent_deep = self.accent_deep or self.accent
        self.accent_light = self.accent_light or self.accent
        self.accent_pale = self.accent_pale or self.line
        self.accent_tint = self.accent_tint or self.bg_card
        self.line_soft = self.line_soft or self.line
        self.charcoal = self.charcoal or self.text_primary
        self.d_bg = self.d_bg or ("16140F" if not self.dark else self.bg_base)
        self.d_panel = self.d_panel or self.d_bg
        self.d_surface = self.d_surface or self.d_panel
        self.d_border = self.d_border or self.line
        self.d_ink = self.d_ink or "FFFFFF"
        self.d_muted = self.d_muted or self.d_ink
        self.d_sub = self.d_sub or self.d_muted
        self.d_dim = self.d_dim or self.d_sub
        self.d_accent_text = self.d_accent_text or self.accent
        self.positive = self.positive or self.accent_sub
        self.negative = self.negative or self.accent
        self.series = self.series or [self.accent, self.accent_sub, self.text_primary, self.text_muted, self.line]


PRESETS = {
    "dark-premium": Theme("dark-premium", "16140F", "1C1912", "FFFFFF", "C9C2B4",
                          "9C937F", "CBA86A", "C0392B", "5A4A2E", dark=True,
                          texture="texture-dark.png"),
    "light-vivid": Theme("light-vivid", "FFFFFF", "F5F9FF", "2C2C2C", "2C2C2C",
                         "7E7E7E", "177DFA", "E9115E", "C9CFD8",
                         texture="texture-light.png"),
}


def load_remember() -> Theme:
    """기본 프리셋 — jc-design-system v2 §6 JSON을 런타임 매핑 (값 하드코딩 금지)."""
    sot = find_sot()
    if sot is None:
        raise FileNotFoundError("jc-design-system 미발견 — ~/.claude/skills 또는 형제 경로에 설치하세요")
    jt = _load_jc_tokens_module(sot)
    tok = jt.load_tokens(sot)
    if not tok:
        raise RuntimeError("signature-tokens.md §6 JSON 파싱 실패")
    c = lambda k: jt.color(tok, k, hash_prefix=False)
    d = lambda k: jt.dark(tok, k, hash_prefix=False)
    fonts = tok.get("font", {}).get("pptxFallback", ["Pretendard"])
    return Theme(
        "remember", c("bg"), c("surface"), c("text"), c("textSecondary"), c("textCaption"),
        c("accent"), c("steel"), c("border"), font_head=fonts[0], font_body=fonts[0],
        surface_alt=c("surfaceAlt"), text_sub=c("textMuted"),
        accent_deep=c("accentStrong"), accent_light=c("accentLight"), accent_pale=c("accentSoftLine"),
        accent_tint=c("accentSoft"), line_soft=c("surfaceSoft"), charcoal=c("primarySoft"),
        d_bg=d("bg"), d_panel=d("panel"), d_surface=d("surface"), d_border=d("border"),
        d_ink=d("text"), d_muted=d("textMuted"), d_sub=d("textSub"), d_dim=d("textDim"),
        d_accent_text=d("accentText"), positive=c("success"), negative=c("danger"),
        series=[s.lstrip("#") for s in tok["color"]["data"]],
        logo_light=str(sot / tok["assets"]["logoLight"]), logo_dark=str(sot / tok["assets"]["logoDark"]),
        objet_dir=str(sot / "assets"),
    )


def load_legacy_jc() -> Theme:
    """구 개인 시그니처 — client-overlays.md §3.3 `deck_theme` 블록에서 로드 (명시 요청 시만)."""
    sot = find_sot()
    if sot is None:
        raise FileNotFoundError("jc-design-system 미발견")
    md = (sot / "references" / "client-overlays.md").read_text(encoding="utf-8")
    for block in re.findall(r"```json\s*\n(.*?)\n```", md, re.S):
        try:
            o = json.loads(block)
        except Exception:
            continue
        if o.get("client_id") == "legacy-jc" and "deck_theme" in o:
            t = o["deck_theme"]
            return Theme("legacy-jc", t["bg_base"], t["bg_card"], t["text_primary"], t["text_body"],
                         t["text_muted"], t["accent"], t["accent_sub"], t["line"],
                         font_head=t.get("font_heading", "Pretendard"), font_body=t.get("font_body", "Pretendard"),
                         series=t.get("data", []))
    raise KeyError("legacy-jc deck_theme 블록을 client-overlays.md에서 찾지 못함")


def get_theme(name: str = "remember") -> Theme:
    if name in (None, "", "remember", "default"):
        return load_remember()
    if name == "legacy-jc":
        return load_legacy_jc()
    if name in PRESETS:
        return PRESETS[name]
    raise KeyError(f"알 수 없는 테마: {name} (remember | legacy-jc | {' | '.join(PRESETS)})")


def C(hexstr: str) -> RGBColor:
    """RULE-PPTX-HEX: '#' 제거 6자리."""
    return RGBColor.from_string(hexstr.lstrip("#"))


def set_font_all(run, name: str) -> None:
    """latin + ea + cs 서체를 모두 지정 (한글 렌더링용 a:ea 필수)."""
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", name)
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.Element(qn("a:ea")); latin.addnext(ea)
    ea.set("typeface", name)
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = etree.Element(qn("a:cs")); ea.addnext(cs)
    cs.set("typeface", name)


# ── Deck ─────────────────────────────────────────────────────────
class Deck:
    """리멤버 룩 덱 빌더. runs 스타일 키: '' | accent | deep | light | sub | muted | body | positive | negative | white | 6자리hex"""

    def __init__(self, theme: Theme, doc_name: str, density: str = "SIMPLE",
                 footer_text: str | None = None, client_logo: str | None = None):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.theme, self.doc_name, self.density = theme, doc_name, density
        self.footer_text = footer_text or doc_name
        self.client_logo = client_logo
        self._blank = self.prs.slide_layouts[6]
        self.page = 0

    # ── 색 해석 ───────────────────────────────────────────────
    def _col(self, style, dark: bool) -> str:
        t = self.theme
        if style and re.fullmatch(r"#?[0-9A-Fa-f]{6}", style):
            return style.lstrip("#")
        table = {
            "": t.d_ink if dark else t.text_primary,
            "accent": t.d_accent_text if dark else t.accent,
            "deep": t.d_accent_text if dark else t.accent_deep,
            "light": t.accent_light,
            "sub": t.d_sub if dark else t.text_sub,
            "muted": t.d_dim if dark else t.text_muted,
            "body": t.d_muted if dark else t.text_body,
            "positive": t.positive, "negative": t.negative,
            "white": "FFFFFF", "steel": t.accent_sub,
        }
        return table.get(style, t.d_ink if dark else t.text_primary)

    # ── 슬라이드 ──────────────────────────────────────────────
    def slide(self, dark: bool = False, bg: str | None = None, texture: str | None = None):
        s = self.prs.slides.add_slide(self._blank)
        fill = s.background.fill; fill.solid()
        fill.fore_color.rgb = C(bg or (self.theme.d_bg if dark else self.theme.bg_base))
        s._dark = dark
        tex = texture if texture is not None else self.theme.texture
        if tex:
            p = ASSETS / tex
            if p.exists():
                pic = s.shapes.add_picture(str(p), 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
                s.shapes._spTree.remove(pic._element); s.shapes._spTree.insert(2, pic._element)
        self.page += 1
        return s

    # ── 텍스트 ────────────────────────────────────────────────
    def text(self, s, x, y, w, h, runs, size, bold=False, align=PP_ALIGN.LEFT, font=None,
             spacing=None, anchor=MSO_ANCHOR.TOP, line_spacing=None, italic=False, inset=0.0,
             para_space_after=None):
        """runs: str | [(text, style)] | [(text, style, bold, size)]. '\\n'은 문단 구분."""
        dark = getattr(s, "_dark", False)
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(inset)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        if isinstance(runs, str):
            runs = [(runs, "")]
        para = tf.paragraphs[0]; para.alignment = align
        if line_spacing: para.line_spacing = line_spacing
        if para_space_after is not None: para.space_after = Pt(para_space_after)
        for item in runs:
            txt, style = item[0], (item[1] if len(item) > 1 else "")
            b = item[2] if len(item) > 2 and item[2] is not None else bold
            sz = item[3] if len(item) > 3 and item[3] is not None else size
            for pi, part in enumerate(str(txt).split("\n")):
                if pi > 0:
                    para = tf.add_paragraph(); para.alignment = align
                    if line_spacing: para.line_spacing = line_spacing
                    if para_space_after is not None: para.space_after = Pt(para_space_after)
                if part == "":
                    continue
                r = para.add_run(); r.text = part
                f = r.font; f.size = Pt(sz); f.bold = b; f.italic = italic
                f.color.rgb = C(self._col(style, dark))
                set_font_all(r, font or self.theme.font_body)
                if spacing is not None:
                    r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
        return tb

    # ── 도형 프리미티브 ────────────────────────────────────────
    def box(self, s, x, y, w, h, fill=None, line_color=None, line_w=0.75, radius=0.06, shape=None, round_=None):
        if round_ is False: radius = 0
        shp = s.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        if radius and shape is None:
            try: shp.adjustments[0] = radius
            except Exception: pass
        if fill: shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
        else: shp.fill.background()
        if line_color: shp.line.color.rgb = C(line_color); shp.line.width = Pt(line_w)
        else: shp.line.fill.background()
        shp.shadow.inherit = False
        if shp.has_text_frame: shp.text_frame.text = ""
        return shp

    def grad(self, s, x, y, w, h, c1=None, c2=None, angle=135, radius=0.06, shape=None):
        """그라디언트 면 — 슬라이드당 1회."""
        shp = s.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        if radius and shape is None:
            try: shp.adjustments[0] = radius
            except Exception: pass
        f = shp.fill; f.gradient(); f.gradient_angle = angle
        st = f.gradient_stops
        st[0].color.rgb = C(c1 or self.theme.accent); st[0].position = 0.0
        st[1].color.rgb = C(c2 or self.theme.accent_light); st[1].position = 1.0
        shp.line.fill.background(); shp.shadow.inherit = False
        return shp

    def hline(self, s, x, y, w, color=None, weight=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = C(color or self.theme.line); ln.line.width = Pt(weight); return ln

    def vline(self, s, x, y, h, color=None, weight=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
        ln.line.color.rgb = C(color or self.theme.line); ln.line.width = Pt(weight); return ln

    def arrow(self, s, x, y, w=0.5, h=0.32, color=None):
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
        ar.fill.solid(); ar.fill.fore_color.rgb = C(color or self.theme.accent)
        ar.line.fill.background(); ar.shadow.inherit = False; return ar

    def chevron(self, s, x, y, w=0.28, h=0.5, color=None):
        ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
        ar.fill.solid(); ar.fill.fore_color.rgb = C(color or self.theme.accent_pale)
        ar.line.fill.background(); ar.shadow.inherit = False; return ar

    def ring(self, s, cx, cy, r, thick=0.22, color=None, gradient=True):
        shp = s.shapes.add_shape(MSO_SHAPE.DONUT, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        try: shp.adjustments[0] = thick / (2 * r)
        except Exception: pass
        if gradient:
            f = shp.fill; f.gradient(); f.gradient_angle = 135
            f.gradient_stops[0].color.rgb = C(self.theme.accent); f.gradient_stops[0].position = 0
            f.gradient_stops[1].color.rgb = C(self.theme.accent_light); f.gradient_stops[1].position = 1
        else:
            shp.fill.solid(); shp.fill.fore_color.rgb = C(color or self.theme.accent_tint)
        shp.line.fill.background(); shp.shadow.inherit = False; return shp

    def dot(self, s, cx, cy, r=0.08, color=None):
        shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        shp.fill.solid(); shp.fill.fore_color.rgb = C(color or self.theme.accent)
        shp.line.fill.background(); shp.shadow.inherit = False; return shp

    def pill(self, s, x, y, w, h, label, fill=None, color=None, size=SZ_LABEL, bold=True):
        shp = self.box(s, x, y, w, h, fill=fill or self.theme.accent_tint, radius=0.5)
        self.text(s, x, y, w, h, [(label, color or self.theme.accent_deep)], size, bold=bold,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return shp

    def alpha(self, shp, a):
        """solid fill 투명도 (0~1, 1=불투명)."""
        sf = shp.fill._xPr.find(qn("a:solidFill"))
        if sf is None: return
        el = etree.SubElement(sf[0], qn("a:alpha")); el.set("val", str(int(a * 100000)))

    def run_alpha(self, tb, a):
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                sf = r._r.get_or_add_rPr().find(qn("a:solidFill"))
                if sf is not None:
                    el = etree.SubElement(sf[0], qn("a:alpha")); el.set("val", str(int(a * 100000)))

    # ── 이미지 ────────────────────────────────────────────────
    def picture(self, s, path, x, y, w, h, cover=True, darken=0.0):
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        if cover and Image is not None:
            iw, ih = Image.open(path).size
            box_r, img_r = w / h, iw / ih
            if img_r > box_r:
                keep = box_r / img_r; c = (1 - keep) / 2; pic.crop_left = c; pic.crop_right = c
            else:
                keep = img_r / box_r; c = (1 - keep) / 2; pic.crop_top = c; pic.crop_bottom = c
        if darken > 0:
            ov = self.box(s, x, y, w, h, fill=self.theme.d_bg, radius=0); self.alpha(ov, darken)
        return pic

    def logo(self, s, x, y, h=0.26, dark=None, path=None):
        dark = getattr(s, "_dark", False) if dark is None else dark
        p = path or (self.theme.logo_dark if dark else self.theme.logo_light)
        if not p or not os.path.exists(p):
            return None
        w = h * 6.1
        if Image is not None:
            iw, ih = Image.open(p).size; w = h * iw / ih
        return s.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))

    def objet(self, s, name, x, y, w, h, fade_left=True):
        """다크 슬라이드 히어로 오브제 + 좌측 페이드. name 예: 'objet-03-slit.png'"""
        p = Path(self.theme.objet_dir or "") / name
        if not p.exists():
            return None
        pic = self.picture(s, p, x, y, w, h, cover=True)
        if fade_left:
            # 좌→우 페이드: 다크 배경색 → 투명 (그라디언트 알파)
            shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w * 0.55), Inches(h))
            f = shp.fill; f.gradient(); f.gradient_angle = 0
            f.gradient_stops[0].color.rgb = C(self.theme.d_bg); f.gradient_stops[0].position = 0
            f.gradient_stops[1].color.rgb = C(self.theme.d_bg); f.gradient_stops[1].position = 1
            gs = shp.fill._xPr.findall(".//" + qn("a:gs"))
            if len(gs) >= 2:
                a1 = etree.SubElement(gs[0][0], qn("a:alpha")); a1.set("val", "100000")
                a2 = etree.SubElement(gs[1][0], qn("a:alpha")); a2.set("val", "0")
            shp.line.fill.background(); shp.shadow.inherit = False
        return pic

    def image_placeholder(self, s, x, y, w, h, brief):
        """사진 미제공 시: 플레이스홀더 + 이미지 브리프 (임의 스톡 금지)."""
        self.box(s, x, y, w, h, fill=self.theme.surface_alt, line_color=self.theme.line)
        self.text(s, x + 0.2, y + h / 2 - 0.35, w - 0.4, 0.7, [("[이미지] ", "accent"), (brief, "muted")],
                  SZ_SMALL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ── 네비게이션 ────────────────────────────────────────────
    def nav(self, s, eyebrow, section=""):
        """아이브로우(좌상) + 푸터(좌 행사명 · 중앙 페이지 · 우 로고/섹션)."""
        dark = getattr(s, "_dark", False)
        self.text(s, MARGIN, EYEBROW_Y, CONTENT_W, 0.32, [(eyebrow, "accent")], SZ_EYEBROW, bold=True, spacing=2)
        self.hline(s, MARGIN, FOOTER_Y - 0.08, CONTENT_W, color=self.theme.d_border if dark else self.theme.line_soft, weight=0.5)
        self.text(s, MARGIN, FOOTER_Y, 5.5, 0.3, [(self.footer_text, "muted")], SZ_FOOTER)
        self.text(s, SLIDE_W / 2 - 1, FOOTER_Y, 2, 0.3, [(f"{self.page:02d}", "muted")], SZ_FOOTER, align=PP_ALIGN.CENTER)
        if self.logo(s, SLIDE_W - MARGIN - 1.6, FOOTER_Y + 0.02, h=0.2) is None and section:
            self.text(s, SLIDE_W - MARGIN - 5.5, FOOTER_Y, 5.5, 0.3, [(section, "muted")], SZ_FOOTER, align=PP_ALIGN.RIGHT)
        elif section:
            self.text(s, SLIDE_W - MARGIN - 7.4, FOOTER_Y, 5.6, 0.3, [(section, "muted")], SZ_FOOTER, align=PP_ALIGN.RIGHT)

    def headline(self, s, runs, sub=None, size=None, y=None, h=1.3):
        self.text(s, MARGIN, y if y is not None else HEADLINE_Y, CONTENT_W, h, runs, size or SZ_HEADLINE,
                  bold=True, font=self.theme.font_head, line_spacing=1.12, spacing=-0.5)
        if sub:
            self.text(s, MARGIN, SUB_Y + (0 if y is None else y - HEADLINE_Y), CONTENT_W, 0.6, [(sub, "sub")], SZ_SUB, line_spacing=1.2)

    def content_slide(self, eyebrow, section, headline_runs, sub=None, dark=False):
        s = self.slide(dark=dark); self.nav(s, eyebrow, section); self.headline(s, headline_runs, sub); return s

    def takeaway(self, s, runs):
        dark = getattr(s, "_dark", False)
        self.box(s, MARGIN, TAKEAWAY_Y, CONTENT_W, 0.6, fill=self.theme.d_surface if dark else self.theme.accent_tint, radius=0.25)
        self.text(s, MARGIN + 0.3, TAKEAWAY_Y, CONTENT_W - 0.6, 0.6, runs, 13.5, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ── 컴포넌트 ──────────────────────────────────────────────
    def card(self, s, x, y, w, h, hl=False):
        dark = getattr(s, "_dark", False)
        if dark:
            return self.box(s, x, y, w, h, fill=self.theme.d_surface, line_color=self.theme.d_accent_text if hl else self.theme.d_border, radius=0.08)
        return self.box(s, x, y, w, h, fill=self.theme.bg_card, line_color=self.theme.accent if hl else self.theme.line,
                        line_w=1.25 if hl else 0.75, radius=0.08)

    def kpi(self, s, x, y, w, h, number, unit="", caption_top="", caption_bottom="", solid=False, num_size=None):
        """KPI 카드. solid=그라디언트 면 + 흰 숫자(슬라이드당 1개) / 기본=카드 + 오렌지 숫자."""
        dark = getattr(s, "_dark", False)
        if solid:
            self.grad(s, x, y, w, h, radius=0.08); numc, topc, botc = "white", self.theme.accent_tint, "white"
        else:
            self.card(s, x, y, w, h); numc, topc, botc = "accent", "muted", "body"
        if caption_top:
            self.text(s, x + 0.15, y + 0.14, w - 0.3, 0.3, [(caption_top, topc)], SZ_LABEL + 0.5, bold=True, align=PP_ALIGN.CENTER, spacing=1)
        ns = num_size or SZ_KPI
        self.text(s, x + 0.1, y + h / 2 - ns / 72 * 0.9, w - 0.2, ns / 72 * 1.8,
                  [(str(number), numc, True, ns), ((" " + unit) if unit else "", numc, True, ns * 0.42)],
                  ns, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=self.theme.font_head)
        if caption_bottom:
            self.text(s, x + 0.15, y + h - 0.62, w - 0.3, 0.55, [(caption_bottom, botc)], SZ_LABEL + 1, align=PP_ALIGN.CENTER, line_spacing=1.15)

    def kpi_card(self, s, x, y, w, h, number, unit="", caption_top="", caption_bottom="", solid=False):
        return self.kpi(s, x, y, w, h, number, unit, caption_top, caption_bottom, solid)

    def label_value_rows(self, s, x, y, w, rows, row_h=0.52):
        for i, (lab, val) in enumerate(rows):
            yy = y + i * row_h
            self.text(s, x, yy, w * 0.28, row_h, [(lab, "deep")], SZ_BODY, bold=True)
            self.text(s, x + w * 0.30, yy, w * 0.70, row_h, [(val, "body")], SZ_BODY)
            self.hline(s, x, yy + row_h - 0.02, w, color=self.theme.line_soft, weight=0.5)

    def numbered_cards(self, s, items, y=None, h=2.1, highlight=None, num_prefix="", body_size=None):
        """items=[(번호, 제목, 설명)] 균등 그리드. highlight=강조 인덱스(상단 그라디언트 룰)."""
        y = CONTENT_Y + 0.5 if y is None else y
        n = len(items); w = (CONTENT_W - GAP * (n - 1)) / n
        for i, (num, title, desc) in enumerate(items):
            x = MARGIN + i * (w + GAP); hl = (highlight is not None and i == highlight)
            self.card(s, x, y, w, h, hl=hl)
            if hl: self.grad(s, x, y, w, 0.09, radius=0)
            self.text(s, x + 0.2, y + 0.2, w - 0.4, 0.35, [(f"{num_prefix}{num}", "accent")], 13 if hl else 12, bold=True, spacing=1)
            self.text(s, x + 0.2, y + 0.58, w - 0.4, 0.62, [(title, "")], SZ_CARD_TITLE + (1 if hl else 0), bold=True, line_spacing=1.1)
            self.text(s, x + 0.2, y + 1.22, w - 0.4, h - 1.35, [(desc, "body")], body_size or (SZ_LABEL + 1.8), line_spacing=1.3)

    def compare_frame(self, s, left, right, y=None, h=2.6):
        """left/right=(태그, 제목, [리스트]). 우측=제안측 강조."""
        y = CONTENT_Y + 0.6 if y is None else y
        w = (CONTENT_W - 1.0) / 2
        for i, ((tag, title, items), hl) in enumerate([(left, False), (right, True)]):
            x = MARGIN + i * (w + 1.0)
            self.card(s, x, y, w, h, hl=hl)
            if hl: self.grad(s, x, y, w, 0.09, radius=0)
            self.text(s, x + 0.25, y + 0.2, w - 0.5, 0.3, [(tag, "accent" if hl else "muted")], SZ_LABEL + 1, bold=True, spacing=1.5)
            self.text(s, x + 0.25, y + 0.55, w - 0.5, 0.45, [(title, "")], SZ_CARD_TITLE, bold=True)
            for j, it in enumerate(items):
                self.dot(s, x + 0.32, y + 1.22 + j * 0.42, r=0.045, color=self.theme.accent if hl else self.theme.line)
                self.text(s, x + 0.47, y + 1.08 + j * 0.42, w - 0.7, 0.42, [(it, "body")], SZ_BODY)
        self.arrow(s, MARGIN + w + 0.22, y + h / 2 - 0.18, 0.56, 0.36)

    def step_pills(self, s, phases, active, y=None):
        y = CONTENT_Y + 0.55 if y is None else y
        n = len(phases); w = (CONTENT_W - GAP * (n - 1)) / n
        for i, (lab, sub) in enumerate(phases):
            x = MARGIN + i * (w + GAP); on = i == active
            self.box(s, x, y, w, 0.46, fill=self.theme.accent_tint if on else None,
                     line_color=self.theme.accent if on else self.theme.line, radius=0.5)
            self.text(s, x, y + 0.06, w, 0.34, [(lab + "  ", "deep" if on else "muted"), (sub, "muted")], SZ_BODY, bold=True, align=PP_ALIGN.CENTER)

    def step_cards(self, s, steps, y=None, h=2.0, highlight=None):
        """T11 프로세스: steps=[(태그, 제목, 설명)] + 셰브론 연결."""
        y = CONTENT_Y + 0.5 if y is None else y
        n = len(steps); w = (CONTENT_W - GAP * (n - 1)) / n
        for i, (tag, title, desc) in enumerate(steps):
            x = MARGIN + i * (w + GAP); hl = (highlight is not None and i == highlight)
            self.card(s, x, y, w, h, hl=hl)
            self.text(s, x + 0.2, y + 0.15, w - 0.4, 0.3, [(tag, "accent")], SZ_LABEL + 1, bold=True, spacing=1.5)
            self.text(s, x + 0.2, y + 0.47, w - 0.4, 0.45, [(title, "")], 14.5, bold=True)
            self.text(s, x + 0.2, y + 0.95, w - 0.4, h - 1.05, [(desc, "body")], SZ_LABEL + 1, line_spacing=1.35)
            if i < n - 1: self.chevron(s, x + w + 0.03, y + h / 2 - 0.25, w=0.27, h=0.5)

    def bullet_list(self, s, x, y, w, items, size=SZ_BODY, gap=0.36, style="body", bullet_color=None):
        for i, it in enumerate(items):
            self.dot(s, x + 0.07, y + i * gap + 0.13, r=0.045, color=bullet_color or self.theme.accent)
            self.text(s, x + 0.22, y + i * gap, w - 0.22, gap, [(it, style)], size, line_spacing=1.15)

    def table(self, s, x, y, w, headers, rows, col_w=None, row_h=0.42, size=SZ_BODY, highlight_rows=(), accent_col=None):
        """DS 표: 가로선만. 헤더 캡션색 + 1.5px 잉크 하단선, 행 1px 보더. 첫 열 600."""
        dark = getattr(s, "_dark", False)
        n = len(headers); col_w = col_w or [w / n] * n
        cx = x
        for j, hd in enumerate(headers):
            self.text(s, cx + 0.08, y, col_w[j] - 0.16, row_h, [(hd, "accent" if j == accent_col else "muted")], size - 2, bold=True, spacing=0.6,
                      anchor=MSO_ANCHOR.MIDDLE); cx += col_w[j]
        self.hline(s, x, y + row_h, w, color=self.theme.d_ink if dark else self.theme.text_primary, weight=1.5)
        for i, row in enumerate(rows):
            yy = y + row_h * (i + 1)
            if i in highlight_rows:
                self.box(s, x, yy + 0.02, w, row_h - 0.04, fill=self.theme.d_surface if dark else self.theme.accent_tint, radius=0.1)
            cx = x
            for j, cell in enumerate(row):
                st = "" if j == 0 else ("deep" if j == accent_col else "body")
                self.text(s, cx + 0.08, yy, col_w[j] - 0.16, row_h, [(str(cell), st)], size, bold=(j == 0), anchor=MSO_ANCHOR.MIDDLE,
                          align=PP_ALIGN.RIGHT if _is_num(cell) else PP_ALIGN.LEFT); cx += col_w[j]
            self.hline(s, x, yy + row_h, w, color=self.theme.d_border if dark else self.theme.line, weight=0.5)

    def profile_capsules(self, s, people, y=None, h=2.6, photos=None):
        """DS 08 인력: people=[(이름, 역할, 담당)] 4열까지. photos=[경로|None]."""
        y = CONTENT_Y + 0.5 if y is None else y
        dark = getattr(s, "_dark", False)
        n = len(people); w = (CONTENT_W - GAP * (n - 1)) / n
        for i, (name, role, duty) in enumerate(people):
            x = MARGIN + i * (w + GAP)
            self.box(s, x, y, w, h, fill=self.theme.charcoal if dark else self.theme.bg_card, line_color=None if dark else self.theme.line, radius=0.08)
            r = 0.55; cx, cy = x + w / 2, y + 0.75
            p = photos[i] if photos and i < len(photos) else None
            if p and os.path.exists(str(p)):
                self.picture(s, p, cx - r, cy - r, 2 * r, 2 * r, cover=True)
            else:
                self.ring(s, cx, cy, r, thick=0.06, gradient=False, color=self.theme.accent_tint)
                self.text(s, cx - r, cy - 0.25, 2 * r, 0.5, [(name[:1], "accent")], 18, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            self.text(s, x + 0.15, y + 1.45, w - 0.3, 0.4, [(name, "white" if dark else "")], SZ_CARD_TITLE, bold=True, align=PP_ALIGN.CENTER)
            self.text(s, x + 0.15, y + 1.85, w - 0.3, 0.3, [(role, "light" if dark else "deep")], SZ_SMALL, bold=True, align=PP_ALIGN.CENTER)
            self.text(s, x + 0.15, y + 2.15, w - 0.3, h - 2.2, [(duty, self.theme.d_muted if dark else "body")], SZ_LABEL + 1, align=PP_ALIGN.CENTER, line_spacing=1.25)

    # ── 표지 · 목차 · 섹션 · 클로징 ─────────────────────────
    def cover(self, doc_label, title, slogan="", date="", issuer="(주)리멤버앤컴퍼니 마켓데이터사업실 · MICE 비즈팀",
              dark=True, objet_name="objet-03-slit.png", cover_image=None):
        """DS 01 표지(다크) / T1 링 표지(라이트). client_logo가 있으면 좌상단 슬롯."""
        s = self.slide(dark=dark)
        if dark:
            if cover_image and os.path.exists(str(cover_image)):
                self.picture(s, cover_image, SLIDE_W - 4.3, 0, 4.3, SLIDE_H, cover=True, darken=0.35)
            else:
                self.objet(s, objet_name, SLIDE_W - 4.3, 0, 4.3, SLIDE_H)
            for yy in (0.85, 6.75):
                ln = self.hline(s, MARGIN, yy, CONTENT_W, color="FFFFFF", weight=1.0)
            if self.client_logo and os.path.exists(self.client_logo):
                self.logo(s, MARGIN, 1.05, h=0.5, path=self.client_logo)
            if date:
                self.text(s, SLIDE_W - MARGIN - 4, 1.12, 4, 0.3, [(f"PROPOSAL   ·   {date}", "light")], 10, bold=True, spacing=2, align=PP_ALIGN.RIGHT)
            self.text(s, MARGIN, 2.3, CONTENT_W, 0.4, [(doc_label, "accent")], 14, bold=True, spacing=2)
            self.text(s, MARGIN, 2.75, 8.4, 2.2, [(title, "")], SZ_COVER, bold=True, font=self.theme.font_head, line_spacing=1.12)
            if slogan:
                self.text(s, MARGIN, 4.85, 8.4, 0.6, [(slogan, "body")], 16)
            self.text(s, MARGIN, 6.88, 8, 0.3, [(issuer, "body")], 10)
            self.logo(s, SLIDE_W - MARGIN - 1.7, 6.9, h=0.18)
        else:
            self.ring(s, SLIDE_W - 1.0, SLIDE_H - 0.6, 2.4, thick=0.26)
            self.dot(s, SLIDE_W - 3.9, SLIDE_H - 1.6, r=0.09)
            self.logo(s, MARGIN, 0.6, h=0.26)
            self.text(s, MARGIN, 1.5, CONTENT_W, 0.35, [(doc_label, "accent")], SZ_EYEBROW, bold=True, spacing=3)
            self.text(s, MARGIN, 1.95, 8.6, 2.0, [(title, "")], 34, bold=True, font=self.theme.font_head, line_spacing=1.12)
            if slogan: self.text(s, MARGIN, 4.1, 8.6, 0.6, [(slogan, "body")], 16)
            meta = "   ·   ".join(x for x in (issuer, date) if x)
            if meta: self.text(s, MARGIN, 6.6, CONTENT_W, 0.4, [(meta, "muted")], 10)
        return s

    def contents(self, items, eyebrow="CONTENTS", title="목차"):
        """DS 02: items=[(번호, 항목)]."""
        s = self.slide(); self.nav(s, eyebrow)
        self.text(s, MARGIN, HEADLINE_Y, 6, 0.9, [(title, "")], SZ_HEADLINE, bold=True, font=self.theme.font_head)
        for i, (num, label) in enumerate(items):
            yy = CONTENT_Y + i * 0.58
            self.text(s, MARGIN, yy, 0.9, 0.5, [(str(num), "accent")], 14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            self.text(s, MARGIN + 0.9, yy, 8, 0.5, [(label, "")], 16, anchor=MSO_ANCHOR.MIDDLE)
            self.hline(s, MARGIN, yy + 0.55, 9.5, color=self.theme.line_soft, weight=0.5)
        self.dot(s, SLIDE_W - 2.2, 3.4, r=0.5)
        return s

    def section_divider(self, num, title, summary="", section_label="", style="dark", objet_name=None, photo=None):
        """DS 03(다크) 또는 T5(그라디언트)."""
        if style == "gradient":
            s = self.slide(dark=True, bg=self.theme.accent)
            self.grad(s, 0, 0, SLIDE_W, SLIDE_H, c1="E8641F", c2=self.theme.accent_light, radius=0)
            if photo and os.path.exists(str(photo)):
                self.picture(s, photo, SLIDE_W * 0.55, 0, SLIDE_W * 0.45, SLIDE_H, darken=0.35)
            rg = self.ring(s, SLIDE_W - 1.2, 0.9, 1.9, thick=0.26, gradient=False, color="FFFFFF"); self.alpha(rg, 0.25)
            self.logo(s, MARGIN, 0.6, h=0.28, dark=True)
            tb = self.text(s, MARGIN, 2.2, 6, 1.4, [(f"{num:02d}", "white")], 88, bold=True, spacing=-2); self.run_alpha(tb, 0.55)
            self.text(s, MARGIN, 3.7, 8.5, 1.2, [(title, "white")], 36, bold=True, line_spacing=1.1, font=self.theme.font_head)
            self.hline(s, MARGIN, 5.05, 3.0, color="FFFFFF", weight=1.25)
            if summary: self.text(s, MARGIN, 5.2, 8.5, 1.0, [(summary, "white")], 14, line_spacing=1.35)
            self.text(s, MARGIN, FOOTER_Y, 7, 0.3, [(self.footer_text, self.theme.accent_tint)], SZ_FOOTER)
            return s
        s = self.slide(dark=True)
        if objet_name:
            self.objet(s, objet_name, SLIDE_W - 4.3, 0, 4.3, SLIDE_H)
        tb = self.text(s, MARGIN, 1.2, 7, 2.8, [(f"{num:02d}", "")], 170, bold=True, font=self.theme.font_head); self.run_alpha(tb, 0.12)
        self.text(s, MARGIN, 4.15, 8, 0.35, [(f"SECTION {num:02d}", "accent")], SZ_EYEBROW, bold=True, spacing=2)
        self.text(s, MARGIN, 4.55, 9.0, 1.2, [(title, "")], 34, bold=True, font=self.theme.font_head)
        if summary: self.text(s, MARGIN, 5.75, 8.5, 0.8, [(summary, "body")], 14, line_spacing=1.35)
        self.text(s, MARGIN, FOOTER_Y, 7, 0.3, [(self.footer_text, "muted")], SZ_FOOTER)
        self.logo(s, SLIDE_W - MARGIN - 1.7, FOOTER_Y + 0.02, h=0.2)
        return s

    def closing(self, title="감사합니다", contact=(), objet_name="objet-06-coil.png", cta=""):
        """DS 10 클로징(다크). contact=[줄...]. cta가 있으면 제목 아래 행동 요청."""
        s = self.slide(dark=True)
        self.objet(s, objet_name, SLIDE_W - 5.2, 0, 5.2, SLIDE_H)
        self.hline(s, MARGIN, 1.6, 4.0, color="FFFFFF", weight=1.0)
        self.text(s, MARGIN, 2.0, 7.5, 1.2, [(title, "")], 30, bold=True, font=self.theme.font_head)
        if cta: self.text(s, MARGIN, 3.15, 7.5, 0.9, [(cta, "light")], 14, line_spacing=1.3)
        yy = 4.3
        for line in contact:
            self.text(s, MARGIN, yy, 7.5, 0.35, [(line, "sub")], 11); yy += 0.36
        self.logo(s, MARGIN, 6.5, h=0.28)
        return s

    def save(self, path):
        self.prs.save(str(path)); return path


def _is_num(v) -> bool:
    return bool(re.fullmatch(r"[\d,.\-+%원명분개회]+", str(v).strip())) and any(ch.isdigit() for ch in str(v))


# ── 스모크 테스트 ────────────────────────────────────────────────
def _smoke(out="deck_kit_smoke.pptx"):
    d = Deck(get_theme("remember"), "스모크 테스트 덱", footer_text="Remember MICE · Smoke")
    d.cover("행사 기획·운영 제안서", "발주처\n행사명 2026", slogan="콘셉트 슬로건", date="2026. 09. 21")
    d.contents([("01", "행사 이해"), ("02", "운영 계획"), ("03", "견적 요약")])
    d.section_divider(1, "행사 이해", "250명, 한 테이블.", objet_name="objet-01-bulb.png")
    s = d.content_slide("01 · 행사 이해", "행사 이해", [("감사가 아니라, ", ""), ("'내년의 첫 미팅'", "accent")], sub="— 관계 195분")
    d.numbered_cards(s, [("01", "타깃", "의사결정권자만 골라내는 설계"), ("02", "쇼업", "등록자를 참석자로"), ("03", "세일즈", "대화를 계약으로")], highlight=2)
    d.takeaway(s, [("좌석은 ", ""), ("의사결정 그룹", "deep"), ("에게만 씁니다.", "")])
    s = d.content_slide("02 · 운영 계획", "운영 계획", [("현장 17명, ", ""), ("지휘선 하나", "accent"), (".", "")])
    d.kpi(s, MARGIN, 2.8, 3.7, 2.2, "250", "명", caption_top="TARGET", caption_bottom="초청 확정", solid=True)
    d.kpi(s, MARGIN + 4.0, 2.8, 3.7, 2.2, "75", "%", caption_top="SHOW-UP", caption_bottom="쇼업률 목표")
    d.table(s, MARGIN + 8.1, 2.8, 4.0, ["항목", "금액"], [["베뉴", "12,000,000"], ["시스템", "8,500,000"], ["운영", "6,200,000"]], highlight_rows=[0], accent_col=1)
    s = d.content_slide("03 · 팀", "투입 인력", [("현장 지휘는 ", ""), ("한 사람", "accent"), ("이 맡습니다.", "")])
    d.profile_capsules(s, [("이OO", "총괄 PM", "발주처 커뮤니케이션·의사결정"), ("김OO", "운영", "베뉴·F&B·동선"), ("박OO", "등록·현장", "등록 데스크·안내")])
    d.closing(contact=["(주)리멤버앤컴퍼니 마켓데이터사업실 · MICE 비즈팀", "mice_solution@remember.co.kr"], cta="다음 단계: 킥오프 미팅 일정을 제안드립니다.")
    return d.save(out)


if __name__ == "__main__":
    print("saved:", _smoke(sys.argv[1] if len(sys.argv) > 1 else "deck_kit_smoke.pptx"))
