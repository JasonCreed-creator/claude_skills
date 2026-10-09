# -*- coding: utf-8 -*-
"""remember_kit — jc-pptx deck_kit 위에 '리멤버 웜 페이퍼 룩' 토큰을 주입한 빌드 헬퍼.

토큰 정본: jc-design-system/references/signature-tokens.md §6 (paper/ink/orange 계열, 구 jc-remember-html 폐합)
지오메트리 정본: jc-pptx/references/slide-types.md (13.33 x 7.5 in 그리드)
"""
import os, sys, io
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ── 그리드 상수 (jc-pptx deck_kit 준용) ─────────────────────────
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.62
CONTENT_W = SLIDE_W - 2 * MARGIN
EYEBROW_Y, HEADLINE_Y, SUB_Y = 0.50, 0.90, 2.05
CONTENT_Y, CONTENT_BOTTOM = 2.50, 6.30
TAKEAWAY_Y, FOOTER_Y = 6.34, 7.06
GAP = 0.33

SZ_EYEBROW, SZ_HEADLINE, SZ_SUB = 11.5, 30, 14
SZ_CARD_TITLE, SZ_BODY, SZ_LABEL, SZ_FOOTER = 15, 11, 8.5, 8
SZ_KPI = 48

# ── 리멤버 웜 페이퍼 룩 토큰 ───────────────────────────────────
T = dict(
    paper="FBFAF6", surface="FFFFFF", surface_warm="F4F1EA",
    ink="1A1A1A", ink_sub="6E6E6E", warm_gray="8C867A",
    border="DCD6C8", line="C9C9C0", line_soft="EFEBE2",
    brown="4A463F", charcoal="332F29",
    orange="EB6F2A", orange_deep="B8431A", orange_soft="F5A05A",
    orange_tint="FFF1E6", orange_pale="F3B48A",
    steel="476580", steel_tint="E8EEF3",
    positive="196B24", negative="D93636",
    # dark mapping
    d_canvas="211E1A", d_objet="141210", d_surface="2A2620", d_border="3E3931",
    d_ink="F4F0E9", d_sub="A89F92", d_dim="6E655A", d_orange="F08A4C", d_steel="8FAEC7",
    white="FFFFFF",
)
FONT = "Pretendard"

def C(h):
    return RGBColor.from_string(h.lstrip("#"))

def _set_font_all(run, name):
    """latin + ea + cs 서체를 모두 지정 (한글 렌더링용 a:ea 필수)."""
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", name)
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.Element(qn("a:ea")); latin.addnext(ea)
    ea.set("typeface", name)
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = etree.Element(qn("a:cs")); ea.addnext(cs)
    cs.set("typeface", name)


class RDeck:
    """리멤버 룩 덱 빌더."""
    def __init__(self, doc_name, assets_dir, img_dir):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.doc_name = doc_name
        self.assets = assets_dir
        self.img = img_dir
        self._blank = self.prs.slide_layouts[6]
        self.page = 0

    # ── 슬라이드 ────────────────────────────────────────────────
    def slide(self, dark=False, bg=None):
        s = self.prs.slides.add_slide(self._blank)
        fill = s.background.fill; fill.solid()
        fill.fore_color.rgb = C(bg or (T["d_canvas"] if dark else T["paper"]))
        s._dark = dark
        self.page += 1
        return s

    # ── 텍스트 프리미티브 ─────────────────────────────────────────
    def text(self, s, x, y, w, h, runs, size, bold=False, align=PP_ALIGN.LEFT,
             color=None, spacing=None, anchor=MSO_ANCHOR.TOP, line_spacing=None,
             italic=False, wrap=True, inset=0.0, para_space_after=None):
        """runs: str | [(text, color_hex_or_None, bold_override_or_None)] | [(text, color)]
        여러 문단은 텍스트 안의 '\\n' 으로 구분."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(inset)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        if isinstance(runs, str):
            runs = [(runs, None)]
        default = color or (T["d_ink"] if getattr(s, "_dark", False) else T["ink"])
        # 문단 분할: 각 run 텍스트 내 \n 을 기준으로 새 문단
        para = tf.paragraphs[0]; para.alignment = align
        if line_spacing: para.line_spacing = line_spacing
        if para_space_after is not None: para.space_after = Pt(para_space_after)
        first = True
        for item in runs:
            txt, col = item[0], item[1]
            b = item[2] if len(item) > 2 and item[2] is not None else bold
            sz = item[3] if len(item) > 3 and item[3] is not None else size
            parts = txt.split("\n")
            for pi, part in enumerate(parts):
                if pi > 0:
                    para = tf.add_paragraph(); para.alignment = align
                    if line_spacing: para.line_spacing = line_spacing
                    if para_space_after is not None: para.space_after = Pt(para_space_after)
                if part == "":
                    continue
                r = para.add_run(); r.text = part
                f = r.font; f.size = Pt(sz); f.bold = b; f.italic = italic
                f.color.rgb = C(col or default)
                _set_font_all(r, FONT)
                if spacing is not None:
                    r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
        return tb

    # ── 도형 ────────────────────────────────────────────────────
    def box(self, s, x, y, w, h, fill=None, line=None, line_w=0.75, radius=0.06,
            shape=None, shadow=False):
        shp = s.shapes.add_shape(
            shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
            Inches(x), Inches(y), Inches(w), Inches(h))
        if radius and shape is None:
            try: shp.adjustments[0] = radius
            except Exception: pass
        if fill: shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
        else: shp.fill.background()
        if line: shp.line.color.rgb = C(line); shp.line.width = Pt(line_w)
        else: shp.line.fill.background()
        shp.shadow.inherit = shadow
        if shp.has_text_frame:
            shp.text_frame.text = ""
        return shp

    def grad(self, s, x, y, w, h, c1=None, c2=None, angle=135, radius=0.06, shape=None):
        """그라디언트 면 (슬라이드당 1회 원칙)."""
        shp = s.shapes.add_shape(
            shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
            Inches(x), Inches(y), Inches(w), Inches(h))
        if radius and shape is None:
            try: shp.adjustments[0] = radius
            except Exception: pass
        f = shp.fill; f.gradient(); f.gradient_angle = angle
        st = f.gradient_stops
        st[0].color.rgb = C(c1 or T["orange"]); st[0].position = 0.0
        st[1].color.rgb = C(c2 or T["orange_soft"]); st[1].position = 1.0
        shp.line.fill.background(); shp.shadow.inherit = False
        return shp

    def hline(self, s, x, y, w, color=None, weight=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                    Inches(x + w), Inches(y))
        ln.line.color.rgb = C(color or T["border"]); ln.line.width = Pt(weight)
        return ln

    def vline(self, s, x, y, h, color=None, weight=0.75):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                    Inches(x), Inches(y + h))
        ln.line.color.rgb = C(color or T["border"]); ln.line.width = Pt(weight)
        return ln

    def arrow(self, s, x, y, w=0.5, h=0.32, color=None):
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
        ar.fill.solid(); ar.fill.fore_color.rgb = C(color or T["orange"])
        ar.line.fill.background(); ar.shadow.inherit = False
        return ar

    def chevron(self, s, x, y, w=0.28, h=0.5, color=None):
        ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
        ar.fill.solid(); ar.fill.fore_color.rgb = C(color or T["orange_pale"])
        ar.line.fill.background(); ar.shadow.inherit = False
        return ar

    def ring(self, s, cx, cy, r, thick=0.22, color=None, grad=True):
        """링 모티프 (도넛). cx,cy 중심, r 반지름 (inch)."""
        shp = s.shapes.add_shape(MSO_SHAPE.DONUT, Inches(cx - r), Inches(cy - r),
                                 Inches(2 * r), Inches(2 * r))
        try: shp.adjustments[0] = thick / (2 * r)
        except Exception: pass
        if grad:
            f = shp.fill; f.gradient(); f.gradient_angle = 135
            f.gradient_stops[0].color.rgb = C(T["orange"]); f.gradient_stops[0].position = 0
            f.gradient_stops[1].color.rgb = C(T["orange_soft"]); f.gradient_stops[1].position = 1
        else:
            shp.fill.solid(); shp.fill.fore_color.rgb = C(color or T["orange_tint"])
        shp.line.fill.background(); shp.shadow.inherit = False
        return shp

    def dot(self, s, cx, cy, r=0.08, color=None):
        shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2*r), Inches(2*r))
        shp.fill.solid(); shp.fill.fore_color.rgb = C(color or T["orange"])
        shp.line.fill.background(); shp.shadow.inherit = False
        return shp

    def pill(self, s, x, y, w, h, label, fill=None, color=None, size=9, bold=True):
        shp = self.box(s, x, y, w, h, fill=fill or T["orange_tint"], radius=0.5)
        self.text(s, x, y, w, h, [(label, color or T["orange_deep"])], size, bold=bold,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return shp

    # ── 이미지 ──────────────────────────────────────────────────
    def picture(self, s, path, x, y, w, h, cover=True, darken=0.0, radius=False):
        """박스에 cover-crop 으로 사진 배치. darken>0 이면 차콜 오버레이."""
        im = Image.open(path); iw, ih = im.size
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        if cover:
            box_r = w / h; img_r = iw / ih
            if img_r > box_r:   # 이미지가 더 넓음 → 좌우 크롭
                keep = box_r / img_r; c = (1 - keep) / 2
                pic.crop_left = c; pic.crop_right = c
            else:
                keep = img_r / box_r; c = (1 - keep) / 2
                pic.crop_top = c; pic.crop_bottom = c
        if darken > 0:
            ov = self.box(s, x, y, w, h, fill=T["d_objet"], radius=0)
            self._alpha(ov, darken)
        return pic

    def logo(self, s, x, y, h=0.26, dark=False):
        fn = "remember-offwhite.png" if dark else "remember-black.png"
        p = os.path.join(self.assets, fn)
        im = Image.open(p); iw, ih = im.size
        w = h * iw / ih
        return s.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))

    def _alpha(self, shp, alpha):
        """solid fill 투명도 (0~1, 1=완전 불투명)."""
        sf = shp.fill._xPr.find(qn("a:solidFill"))
        if sf is None: return
        clr = sf[0]
        a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(int(alpha * 100000)))

    # ── 네비게이션 ──────────────────────────────────────────────
    def nav(self, s, eyebrow, section, dark=None):
        dark = s._dark if dark is None else dark
        acc = T["d_orange"] if dark else T["orange"]
        mut = T["d_dim"] if dark else T["warm_gray"]
        self.text(s, MARGIN, EYEBROW_Y, CONTENT_W, 0.32, [(eyebrow, acc)], SZ_EYEBROW,
                  bold=True, spacing=2)
        self.hline(s, MARGIN, FOOTER_Y - 0.08, CONTENT_W, color=T["d_border"] if dark else T["line_soft"], weight=0.5)
        self.text(s, MARGIN, FOOTER_Y, 7.0, 0.3, [(self.doc_name, mut)], SZ_FOOTER)
        self.text(s, SLIDE_W - MARGIN - 6.0, FOOTER_Y, 6.0, 0.3,
                  [(f"{section}   ·   {self.page:02d}", mut)], SZ_FOOTER, align=PP_ALIGN.RIGHT)

    def headline(self, s, runs, sub=None, size=None, dark=None, y=None, h=1.3):
        dark = s._dark if dark is None else dark
        self.text(s, MARGIN, y if y is not None else HEADLINE_Y, CONTENT_W, h, runs,
                  size or SZ_HEADLINE, bold=True, line_spacing=1.12, spacing=-0.5)
        if sub:
            self.text(s, MARGIN, SUB_Y + (0 if y is None else (y - HEADLINE_Y)), CONTENT_W, 0.6,
                      [(sub, T["d_sub"] if dark else T["ink_sub"])], SZ_SUB, line_spacing=1.2)

    def content_slide(self, eyebrow, section, runs, sub=None, dark=False):
        s = self.slide(dark=dark); self.nav(s, eyebrow, section)
        self.headline(s, runs, sub)
        return s

    def takeaway(self, s, runs, dark=None):
        dark = s._dark if dark is None else dark
        self.box(s, MARGIN, TAKEAWAY_Y, CONTENT_W, 0.6,
                 fill=T["d_surface"] if dark else T["orange_tint"], radius=0.25)
        self.text(s, MARGIN + 0.3, TAKEAWAY_Y, CONTENT_W - 0.6, 0.6, runs, 13.5, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ── 컴포넌트 ────────────────────────────────────────────────
    def card(self, s, x, y, w, h, hl=False, dark=None):
        dark = s._dark if dark is None else dark
        if dark:
            return self.box(s, x, y, w, h, fill=T["d_surface"], line=T["d_orange"] if hl else T["d_border"], radius=0.08)
        return self.box(s, x, y, w, h, fill=T["surface"], line=T["orange"] if hl else T["border"],
                        line_w=1.25 if hl else 0.75, radius=0.08)

    def kpi(self, s, x, y, w, h, number, unit="", top="", bottom="", solid=False, num_size=None, dark=None):
        dark = s._dark if dark is None else dark
        if solid:
            self.grad(s, x, y, w, h, radius=0.08)
            numc, topc, botc = T["white"], T["orange_tint"], T["white"]
        else:
            self.card(s, x, y, w, h)
            numc = T["d_orange"] if dark else T["orange"]
            topc = T["d_sub"] if dark else T["warm_gray"]
            botc = T["d_ink"] if dark else T["brown"]
        if top:
            self.text(s, x + 0.15, y + 0.14, w - 0.3, 0.3, [(top, topc)], SZ_LABEL + 0.5, bold=True,
                      align=PP_ALIGN.CENTER, spacing=1)
        ns = num_size or SZ_KPI
        self.text(s, x + 0.1, y + h / 2 - ns / 72 * 0.9, w - 0.2, ns / 72 * 1.8,
                  [(str(number), numc, True, ns), ((" " + unit) if unit else "", numc, True, ns * 0.42)],
                  ns, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if bottom:
            self.text(s, x + 0.15, y + h - 0.62, w - 0.3, 0.55, [(bottom, botc)], SZ_LABEL + 1,
                      align=PP_ALIGN.CENTER, line_spacing=1.15)

    def numbered_cards(self, s, items, y=None, h=2.1, highlight=None, num_prefix="", body_size=None):
        y = CONTENT_Y + 0.5 if y is None else y
        n = len(items); w = (CONTENT_W - GAP * (n - 1)) / n
        for i, (num, title, desc) in enumerate(items):
            x = MARGIN + i * (w + GAP)
            hl = (highlight is not None and i == highlight)
            self.card(s, x, y, w, h, hl=hl)
            if hl:
                self.grad(s, x, y, w, 0.09, radius=0)
            self.text(s, x + 0.2, y + 0.2, w - 0.4, 0.35, [(f"{num_prefix}{num}", T["orange"])],
                      13 if hl else 12, bold=True, spacing=1)
            self.text(s, x + 0.2, y + 0.58, w - 0.4, 0.62, [(title, None)],
                      SZ_CARD_TITLE + (1 if hl else 0), bold=True, line_spacing=1.1)
            self.text(s, x + 0.2, y + 1.22, w - 0.4, h - 1.35, [(desc, T["brown"])],
                      body_size or (SZ_LABEL + 1.8), line_spacing=1.3)

    def section_divider(self, num, title, summary, section_label, photo=None):
        """T5 · 그라디언트 풀블리드 디바이더 (오브제 없이, 링 모티프)."""
        s = self.slide(dark=True, bg=T["orange"])
        self.grad(s, 0, 0, SLIDE_W, SLIDE_H, c1="E8641F", c2="F5A05A", angle=135, radius=0)
        if photo and os.path.exists(photo):
            self.picture(s, photo, SLIDE_W * 0.55, 0, SLIDE_W * 0.45, SLIDE_H, darken=0.35)
        # 화이트 25% 링 우상단
        rg = self.ring(s, SLIDE_W - 1.2, 0.9, 1.9, thick=0.26, grad=False, color=T["white"])
        self._alpha(rg, 0.25)
        self.logo(s, MARGIN, 0.6, h=0.28, dark=True)
        self.text(s, MARGIN, 2.2, 6, 1.4, [(f"{num:02d}", "FFFFFF")], 88, bold=True, spacing=-2)
        self._set_run_alpha(s.shapes[-1], 0.55)
        self.text(s, MARGIN, 3.7, 8.5, 1.2, [(title, "FFFFFF")], 36, bold=True, line_spacing=1.1)
        self.hline(s, MARGIN, 5.05, 3.0, color="FFFFFF", weight=1.25)
        self.text(s, MARGIN, 5.2, 8.5, 1.0, [(summary, "FFFFFF")], 14, line_spacing=1.35)
        self.text(s, MARGIN, FOOTER_Y, 7, 0.3, [(self.doc_name, "FFF1E6")], SZ_FOOTER)
        self.text(s, SLIDE_W - MARGIN - 6, FOOTER_Y, 6, 0.3, [(f"{section_label}   ·   {self.page:02d}", "FFF1E6")],
                  SZ_FOOTER, align=PP_ALIGN.RIGHT)
        return s

    def _set_run_alpha(self, tb, alpha):
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                rPr = r._r.get_or_add_rPr()
                sf = rPr.find(qn("a:solidFill"))
                if sf is not None:
                    a = etree.SubElement(sf[0], qn("a:alpha")); a.set("val", str(int(alpha * 100000)))

    def bullet_list(self, s, x, y, w, items, size=11, gap=0.36, color=None, bullet_color=None, bold_first=False):
        for i, it in enumerate(items):
            self.dot(s, x + 0.07, y + i * gap + 0.13, r=0.045, color=bullet_color or T["orange"])
            self.text(s, x + 0.22, y + i * gap, w - 0.22, gap, [(it, color or T["brown"])], size, line_spacing=1.15)

    def save(self, path):
        self.prs.save(path); return path
