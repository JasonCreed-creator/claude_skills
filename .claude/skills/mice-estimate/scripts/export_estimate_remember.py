#!/usr/bin/env python3
"""
export_estimate_remember.py — calc_estimate_remember 결과 → 리멤버 견적서(.xlsx) 자동 작성 브리지

레이아웃 정본: 컨피규레이터 src/lib/exportEstimate.js 출력물에서 실측 추출한 스타일 스펙
(다크 타이틀 밴드 · 오렌지 라벨 · 7섹션 · 옵션 O/X 자동 재계산 · PCO 만원 절사 수식).
색상 정본: jc-design-system v2(리멤버 웜 페이퍼) signature-tokens.md §6 JSON — estimate_tokens.palette()로
런타임 로드(SoT 미탐지 시 §6 v2.1.0 폴백). 글꼴 크기·굵기·정렬·테두리·숫자서식·열 너비·행 높이는 실측값 그대로.

재계산: recalc(path) — xlsx 스킬 recalc.py(환경변수 XLSX_RECALC → 형제 경로 → ~/.claude/skills[/synced/*])
→ 없으면 LibreOffice headless 변환 → 둘 다 없으면 Excel에서 열어 저장 안내 후 False.
자가 테스트: python3 export_estimate_remember.py --self-test  (calc → export → recalc → verify, 팔레트 출처 출력)

무결성 게이트: recalc 후 verify(path, result) 로 D10↔pk, D11↔pk_excluding_options 0원 일치 확인.
한글 금액: 검증 환경(LibreOffice)이 NUMBERSTRING 미지원이라 기본은
'정적 한글 + TEXT(D10) 동적 숫자' 하이브리드. Excel 네이티브 연동이 필요하면
use_numberstring=True (단, recalc 검증 불가 — #NAME?).
"""
import os, sys, math, json, shutil, subprocess, tempfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from korean_amount import num_to_korean as _kor_ext  # noqa: E402
from estimate_tokens import palette  # noqa: E402

# jc-design-system v2 §6 토큰을 모듈 로드 시 1회 읽는다. 색 리터럴은 이 dict 밖에 두지 않는다.
# _P['_source'] == 'sot:<경로>' | 'fallback'
_P = palette()


def _kor(n):
    try:
        return _kor_ext(int(n))
    except Exception:
        # NUMBERSTRING 스타일 폴백 (일십/일백/일천 표기)
        digits = ' 일이삼사오육칠팔구'
        small = ['', '십', '백', '천']
        units = ['', '만', '억', '조']

        def four(x):
            s = ''
            for i, ch in enumerate(str(x).zfill(4)):
                d = int(ch)
                if d:
                    s += digits[d].strip() + small[3 - i]
            return s
        parts, i, n = [], 0, int(n)
        while n > 0:
            n, chunk = divmod(n, 10000)
            if chunk:
                parts.append(four(chunk) + units[i])
            i += 1
        return ''.join(reversed(parts)) or '영'


# ---------- 스타일 스펙 (레이아웃: 컨피규레이터 출력물 실측 · 색: jc-design-system v2 §6 역할 키) ----------
THIN = Border(*[Side(style='thin')] * 4)


def _F(size=10, bold=False, color=None):
    return Font(size=size, bold=bold, color=color)


def _fill(hexv):
    return PatternFill('solid', start_color=hexv)


ST = {
    'title':     dict(font=_F(20, True, _P['paper']), fill=_fill(_P['ink']), align=Alignment('center', 'center')),  # 역할: paper 글자 (SoT color.surface) · ink 배경 (SoT color.primary)
    'label_o':   dict(font=_F(10, True, _P['paper']), fill=_fill(_P['accent']), align=Alignment('center', 'center', wrap_text=True), border=THIN),  # 역할: paper 글자 (SoT color.surface) · accent 배경 (SoT color.accent)
    'label_g':   dict(font=_F(10, True), fill=_fill(_P['surfaceAlt']), align=Alignment('center', 'center', wrap_text=True), border=THIN),  # 역할: surfaceAlt 배경 (SoT color.surfaceAlt)
    'val':       dict(font=_F(10), align=Alignment(vertical='center', wrap_text=True), border=THIN),
    'tot_kor':   dict(font=_F(10, True), fill=_fill(_P['accentSoft']), align=Alignment(vertical='center', wrap_text=True), border=THIN),  # 역할: accentSoft 배경 (SoT color.accentSoft)
    'tot_val':   dict(font=_F(12, True, _P['accent']), fill=_fill(_P['accentSoft']), align=Alignment('right', 'center'), border=THIN, numfmt='₩#,##0'),  # 역할: accent 글자 (SoT color.accent) · accentSoft 배경 (SoT color.accentSoft)
    'tot_sub':   dict(font=_F(10, True), fill=_fill(_P['accentSoft']), align=Alignment('center', 'center', wrap_text=True), border=THIN),  # 역할: accentSoft 배경 (SoT color.accentSoft)
    'tot_vat':   dict(font=_F(10, True), fill=_fill(_P['accentSoft']), align=Alignment('right', 'center'), border=THIN, numfmt='₩#,##0'),  # 역할: accentSoft 배경 (SoT color.accentSoft)
    'sec_hdr':   dict(font=_F(10, True), fill=_fill(_P['amberTint']), align=Alignment(vertical='center', wrap_text=True), border=THIN),  # 역할: amberTint 배경 (SoT color.point.amberTint)
    'sec_amt':   dict(font=_F(10, True, _P['paper']), fill=_fill(_P['ink']), align=Alignment('right', 'center'), border=THIN, numfmt='₩#,##0'),  # 역할: paper 글자 (SoT color.surface) · ink 배경 (SoT color.primary)
    'warn':      dict(font=_F(10, True, _P['accentStrong']), fill=_fill(_P['warningBg']), align=Alignment('center', 'center', wrap_text=True), border=THIN),  # 역할: accentStrong 글자 (SoT color.accentStrong) · warningBg 배경 (SoT color.semantic.warningBg)
    'note_blue': dict(font=_F(10, False, _P['steel']), fill=_fill(_P['steelTint']), align=Alignment('center', 'center', wrap_text=True), border=THIN),  # 역할: steel 글자 (SoT color.point.steel) · steelTint 배경 (SoT color.point.steelTint)
    'col_hdr':   dict(font=_F(10, True, _P['paper']), fill=_fill(_P['inkSoft']), align=Alignment('center', 'center', wrap_text=True), border=THIN),  # 역할: paper 글자 (SoT color.surface) · inkSoft 배경 (SoT color.primarySoft)
    'data':      dict(font=_F(10), align=Alignment(vertical='center', wrap_text=True), border=THIN),
    'data_num':  dict(font=_F(10), align=Alignment('right', 'center', wrap_text=True), border=THIN, numfmt='#,##0'),
    'data_sel':  dict(font=_F(10, True), align=Alignment('center', 'center', wrap_text=True), border=THIN),
    'sub':       dict(font=_F(10, True), fill=_fill(_P['surfaceSoft']), align=Alignment('right', 'center'), border=THIN),  # 역할: surfaceSoft 배경 (SoT color.surfaceSoft)
    'sub_amt':   dict(font=_F(10, True), fill=_fill(_P['surfaceSoft']), align=Alignment('right', 'center'), border=THIN, numfmt='₩#,##0'),  # 역할: surfaceSoft 배경 (SoT color.surfaceSoft)
    'footer':    dict(font=_F(10, False, _P['danger']), align=Alignment('center', 'center'), border=THIN),  # 역할: danger 글자 (SoT color.semantic.danger)
}

WIDTHS = {'A': 22.0, 'B': 26.0, 'C': 51.58, 'D': 16.0, 'E': 10.0, 'F': 18.0, 'G': 28.0, 'H': 11.58}

# ---------- 항목 라벨 (질문이 필요 없는 견적서 — 평이 한국어) ----------
SYS_ROWS = {  # key: (항목, 내용, 산출내역 템플릿 fn(t, amount), 비고 fn)
    'video': ('영상 콘솔', 'LED 디스플레이 운용',
              lambda t, a: '영상 스위처 + 송출 콘솔 + 배분기 + 케이블 일체', lambda t: ''),
    'scaler4k': ('4K 스케일러/KVM', '외부 스케일러+KVM',
                 lambda t, a: '100명 이상 LED 행사 기본 포함', lambda t: '100명 이상 · LED 선택 시 자동 적용'),
    'audio': ('음향 시스템', '하우스 PA + 운용',
              lambda t, a: f'기본 1,500,000원 + 인원 규모 가산 {a-1500000:,}원', lambda t: '50명 초과분 100명 단위로 500,000원 가산'),
    'engineer': ('엔지니어', '영상·음향 운용',
                 lambda t, a: '영상·음향 엔지니어 1인 · 1일 상주', lambda t: ''),
    'presentation': ('발표지원', '프롬프터·모니터·클리커',
                     lambda t, a: '프롬프터 + 발표자 모니터 + 무선 클리커 + 백업 노트북', lambda t: ''),
    'registration': ('등록시스템', '체크인/명찰',
                     lambda t, a: (f'기본 1,000,000원 + 셀프 체크인 키오스크 {t//100}대' if t > 200
                                   else '기본 1,000,000원' + (f' + 초과 인원 가산 {a-1000000:,}원' if a > 1000000 else '')),
                     lambda t: ('키오스크 100명당 1대' if t > 200 else '100명 초과 인원당 5,000원')),
    'misc': ('기타 필요물품', '케이블·소모품 등',
             lambda t, a: f'기본 500,000원 + 인원 규모 가산 {a-500000:,}원', lambda t: '50명 초과분 100명 단위로 500,000원 가산'),
}
DES_ROWS = {
    'env': ('환경조성', '사인·배너·연출', lambda t, a: '무대 배경 + 사인물 + 동선 배너 + 현장 연출 일체', '인원 구간별 정액'),
    'web': ('웹페이지', '마이크로 랜딩', lambda t, a: '행사 소개 페이지 + 온라인 사전등록 폼 + 반응형', ''),
    'kv': ('키비주얼', '메인 + 베리에이션', lambda t, a: '메인 키비주얼 1종 + 배너·현수막·명찰 응용 전개', ''),
}
OPT_ROWS = {  # key: (항목, 내용, 산출내역 fn(result), 비고)
    'souvenir': ('기념품', '참석 기념품', lambda r: f"인당 {r['otBreakdown']['souvenir']//max(1,r['souvenirQty']):,}원 × {r['souvenirQty']}명", ''),
    'emcee': ('사회자', '전문 사회자', lambda r: '기업회의급', ''),
    'photo': ('사진 촬영', '현장 스케치', lambda r: '전문 포토그래퍼 1인 · 보정본 제공', '사진/영상/AVING 택1'),
    'video_sketch': ('영상 스케치', '하이라이트 영상', lambda r: '촬영 + 편집 하이라이트 1편', '사진/영상/AVING 택1'),
    'aving': ('AVING', '미디어 패키지', lambda r: '사진+영상+보도 · AVING Korea/USA', '사진/영상/AVING 택1'),
    'scaler4k': ('4K 스케일러/KVM', '외부 스케일러+KVM', lambda r: '100명 미만 LED 행사 선택 옵션', '100명 미만 전용'),
    'screenRelay': ('화면중계', 'LED 화면 송출 중계', lambda r: '무대 화면 실시간 중계 송출', 'LED 선택 시에만 적용'),
    'fullRecording': ('풀영상 녹화', '전체 세션 녹화', lambda r: '전 세션 풀녹화 + 원본 제공', ''),
    'survey': ('만족도 설문', '사후 설문', lambda r: '설문 설계 + 배포 + 결과 리포트', ''),
    'photowall_basic': ('포토월 (일반형)', '포토존 백월', lambda r: '일반형 포토월 제작·설치', '일반형/고급형 택1'),
    'photowall_premium': ('포토월 (고급형)', '포토존 백월', lambda r: '고급형 포토월 제작·설치', '일반형/고급형 택1'),
    'rsvpHandling': ('참가확정 응대 대행', '순수 응대 대행', lambda r: f"참석 {r['t']}명 × 인당 20,000원", '모객 게런티 아님'),
    'booth_standard': ('부스 설치 (일반형)', '조립식 부스 (설치·철거 포함)', None, ''),
    'booth_premium': ('부스 설치 (프리미엄)', '프리미엄 부스 (설치·철거 포함)', None, ''),
}


def _apply(cell, key):
    s = ST[key]
    cell.font = s['font']
    if 'fill' in s:
        cell.fill = s['fill']
    if 'border' in s:
        cell.border = s['border']
    cell.alignment = s['align']
    if 'numfmt' in s:
        cell.number_format = s['numfmt']


def export_remember_estimate(result, meta, out_path, use_numberstring=False):
    """result: calc_estimate_remember() 반환값 / meta: 외부 주입 슬롯 dict
    meta 키: project_title(필수), venue_text, venue_type, remark, proposal_date,
    validity, supplier_company, supplier_address, supplier_manager, quote_date,
    targeting(모객 C열 텍스트), sheet_name,
    booth_count / booth_premium_count(선택 — 섹션 5 부스 행의 단가×수량 분해용; 없으면 수량 1·단가=금액)"""
    r = result
    if r.get('isCustom'):
        raise ValueError('target 500명 초과 — 자동견적 불가(isCustom), 별도 협의 견적으로 진행')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = meta.get('sheet_name', '리멤버MICE솔루션')
    for c, w in WIDTHS.items():
        ws.column_dimensions[c].width = w

    # ---- 타이틀·메타 블록
    ws.merge_cells('A1:H3')
    ws['A1'] = '리멤버 MICE 솔루션 견적서'
    _apply(ws['A1'], 'title')
    for rr in (1, 2, 3):
        ws.row_dimensions[rr].height = 22.65
    ws.merge_cells('A4:H4')
    ws.row_dimensions[4].height = 4.0

    g_txt = f"{r['t']}명 / {r['g']}명" + (f"  ·  {r['kpi']['text']}" if r.get('kpi') else '')
    left = [('프로젝트명', meta.get('project_title', '')), ('참석/모객', g_txt),
            ('베뉴', meta.get('venue_text', '')), ('비  고', meta.get('remark', '1일 full day 기준')),
            ('견적일시', meta.get('quote_date', ''))]
    right = [('제안일자', meta.get('proposal_date', '')), ('유효기간', meta.get('validity', '제안일자로 부터 30일')),
             ('공 급 자', meta.get('supplier_company', '')), ('주     소', meta.get('supplier_address', '')),
             ('담 당 자', meta.get('supplier_manager', ''))]
    for i, ((la, va), (lb, vb)) in enumerate(zip(left, right)):
        rr = 5 + i
        ws.merge_cells(f'B{rr}:C{rr}')
        ws.merge_cells(f'G{rr}:H{rr}')
        ws[f'A{rr}'] = la; _apply(ws[f'A{rr}'], 'label_o')
        ws[f'B{rr}'] = va; _apply(ws[f'B{rr}'], 'val')
        _apply(ws[f'D{rr}'], 'val'); _apply(ws[f'E{rr}'], 'val')
        ws[f'F{rr}'] = lb; _apply(ws[f'F{rr}'], 'label_g')
        ws[f'G{rr}'] = vb; _apply(ws[f'G{rr}'], 'val')

    # ---- 섹션 본문 (수식은 소계 행번호 확정 후 기입하므로 2-pass)
    row = 13
    sec_sub = {}     # 섹션번호 → 소계 행
    cell_ref = {}    # 'rsvp'/'gen' → 데이터 행

    def sec_header(no, title, note=None):
        nonlocal row
        ws.merge_cells(f'A{row}:B{row}')
        ws.merge_cells(f'D{row}:E{row}')
        ws.merge_cells(f'G{row}:H{row}')
        ws[f'A{row}'] = title
        for cc in 'ABCDE':
            _apply(ws[f'{cc}{row}'], 'sec_hdr')
        _apply(ws[f'F{row}'], 'sec_amt')
        _apply(ws[f'G{row}'], 'sec_hdr')
        hdr_row = row
        row += 1
        if note:
            ws.merge_cells(f'A{row}:H{row}')
            ws[f'A{row}'] = note[0]
            _apply(ws[f'A{row}'], note[1])
            ws.row_dimensions[row].height = 18.65
            row += 1
        return hdr_row

    def col_header(sel=False):
        nonlocal row
        heads = ['항목', '내용', '산출 내역', '단가', '수량', '금액', '선택' if sel else '비고']
        ws.merge_cells(f'G{row}:H{row}')
        for cc, h in zip('ABCDEFG', heads):
            ws[f'{cc}{row}'] = h
            _apply(ws[f'{cc}{row}'], 'col_hdr')
        _apply(ws[f'H{row}'], 'col_hdr')
        row += 1

    def data_row(a, b, c, d, e, g, sel=None):
        nonlocal row
        ws.merge_cells(f'G{row}:H{row}')
        ws[f'A{row}'] = a; _apply(ws[f'A{row}'], 'data')
        ws[f'B{row}'] = b; _apply(ws[f'B{row}'], 'data')
        ws[f'C{row}'] = c; _apply(ws[f'C{row}'], 'data')
        ws[f'D{row}'] = d; _apply(ws[f'D{row}'], 'data_num')
        ws[f'E{row}'] = e; _apply(ws[f'E{row}'], 'data_num')
        if sel is not None:
            ws[f'F{row}'] = f'=D{row}*E{row}*IF(G{row}="O",1,0)'
            ws[f'G{row}'] = sel
            _apply(ws[f'G{row}'], 'data_sel')
        else:
            ws[f'F{row}'] = f'=D{row}*E{row}'
            ws[f'G{row}'] = g
            _apply(ws[f'G{row}'], 'data')
        _apply(ws[f'F{row}'], 'data_num')
        _apply(ws[f'H{row}'], 'data')
        dr = row
        row += 1
        return dr

    def subtotal(first, last):
        nonlocal row
        ws.merge_cells(f'A{row}:E{row}')
        ws.merge_cells(f'G{row}:H{row}')
        ws[f'A{row}'] = '소계'
        for cc in 'ABCDE':
            _apply(ws[f'{cc}{row}'], 'sub')
        ws[f'F{row}'] = f'=SUM(F{first}:F{last})' if last >= first else 0
        _apply(ws[f'F{row}'], 'sub_amt')
        _apply(ws[f'G{row}'], 'sub')
        sr = row
        row += 2  # 소계 + 빈 행
        return sr

    # 1. 베뉴
    h1 = sec_header(1, '1. 베뉴 사용료',
                    ('⚠️ 호텔 비용은 베뉴 협의 대관료 기준이며 확정 금액이 아닙니다. 반드시 베뉴별 견적 확인을 통해 확정해야 합니다.', 'warn'))
    col_header()
    d = data_row('장소 사용료', meta.get('venue_type', '5성급 호텔'), meta.get('venue_text', ''), r['s1'], 1,
                 'F&B 대관료의 80%까지 포함')
    sec_sub[1] = subtotal(d, d)

    # 2. 시스템
    h2 = sec_header(2, '2. 시스템 구축')
    col_header()
    first = row
    for k, amt in r['sysBreakdown'].items():
        if amt <= 0:
            continue
        a, b, cf, gf = SYS_ROWS[k]
        data_row(a, b, cf(r['t'], amt), amt, 1, gf(r['t']))
    sec_sub[2] = subtotal(first, row - 1)

    # 3. 디자인
    h3 = sec_header(3, '3. 디자인·브랜딩')
    col_header()
    first = row
    for k, amt in r['desBreakdown'].items():
        a, b, cf, g = DES_ROWS[k]
        data_row(a, b, cf(r['t'], amt), amt, 1, g)
    sec_sub[3] = subtotal(first, row - 1)

    # 4. 운영
    h4 = sec_header(4, f"4. 운영인력·등록·보험 ({r['t']}명 기준)")
    col_header()
    first = row
    ob = r['opsBreakdown']
    ndesk = math.ceil(r['t'] / 50)
    data_row('등록데스크 지원', '체크인 운영', f"참석 50명당 운영요원 1명 ({r['t']}명 → {ndesk}명)",
             250000, ndesk, '운영요원 1명당 250,000원')
    data_row('행사 운영 인력', '현장 매니저·진행', f"기본 900,000원 + 인원 규모 가산 {ob['ops']-900000:,}원",
             ob['ops'], 1, '100명 초과 시 100명 단위로 300,000원 가산')
    data_row('행사장 안전보험', '행사장 책임보험', f"기본 400,000원 + 인원 규모 가산 {ob['insurance']-400000:,}원",
             ob['insurance'], 1, '100명 초과 시 100명 단위로 100,000원 가산')
    sec_sub[4] = subtotal(first, row - 1)

    # 5. 추가옵션
    h5 = sec_header(5, '5. 추가옵션',
                    ("✅ 오른쪽 '선택' 칸에 O=포함 / X=제외 입력 시 금액·PCO 기획료·총액이 자동 재계산됩니다.", 'note_blue'))
    col_header(sel=True)
    first = row
    for k, amt in r['otBreakdown'].items():
        a, b, cf, g = OPT_ROWS[k]
        if k in ('booth_standard', 'booth_premium'):
            # 단가·수량 분해: otBreakdown 금액 = unit × count. meta로 넘어온 count 우선.
            cnt = meta.get('booth_count' if k == 'booth_standard' else 'booth_premium_count')
            unit = amt // cnt if cnt else amt
            if not cnt:
                cnt, unit = 1, amt
            data_row(a, b, f'부스당 {unit:,}원 × {cnt}개', unit, cnt, '', sel='O')
        elif k == 'souvenir':
            qty = r['souvenirQty'] or 1
            data_row(a, b, cf(r), amt // qty, qty, g, sel='O')
        else:
            data_row(a, b, cf(r), amt, 1, g, sel='O')
    sec_sub[5] = subtotal(first, row - 1)

    # 7섹션 데이터 행이 PCO 수식에 필요하므로, 모객 행을 먼저 확보하기 위해
    # PCO(6) 섹션 수식은 셀 참조 문자열로 뒤에서 채운다 → 행 위치 선확정 방식 유지
    # 6. PCO
    h6 = sec_header(6, '6. PCO 기획료')
    ws.merge_cells(f'G{row}:H{row}')
    for cc, h in zip('ABCDEFG', ['항목', '내용', '산출 내역', '운영비', '요율', '금액', '비고']):
        ws[f'{cc}{row}'] = h
        _apply(ws[f'{cc}{row}'], 'col_hdr')
    _apply(ws[f'H{row}'], 'col_hdr')
    row += 1
    pco_row = row
    ws.merge_cells(f'G{row}:H{row}')
    ws[f'A{row}'] = 'PCO 기획료'; _apply(ws[f'A{row}'], 'data')
    ws[f'B{row}'] = '운영비 × 25%'; _apply(ws[f'B{row}'], 'data')
    ws[f'C{row}'] = '운영비 합계의 25% (섹션 1~5 + 사전신청 관리' + (' + 참관객 응대' if r['genCount'] else '') + ')'
    _apply(ws[f'C{row}'], 'data')
    _apply(ws[f'D{row}'], 'data_num')
    ws[f'E{row}'] = 0.25; _apply(ws[f'E{row}'], 'data_num'); ws[f'E{row}'].number_format = '0%'
    ws[f'F{row}'] = f'=FLOOR(D{row}*E{row},10000)'; _apply(ws[f'F{row}'], 'data_num')
    ws[f'G{row}'] = '만원 미만 절사 · 쇼업 보장 비용 제외'; _apply(ws[f'G{row}'], 'data')
    _apply(ws[f'H{row}'], 'data')
    row += 1
    sec_sub[6] = subtotal(pco_row, pco_row)

    # 7. 모객 솔루션
    h7 = sec_header(7, f"7. 리멤버 모객 솔루션 [게런티 {r['g']}명]" if r['g'] else '7. 리멤버 모객 솔루션 (모객 제외)')
    col_header()
    first = row
    if r['g'] > 0:
        cell_ref['rsvp'] = data_row('사전신청 관리 (RSVP)', '참가확정 관리',
                                    meta.get('targeting', f"게런티 {r['g']}명 × 인당 40,000원"),
                                    40000, r['g'], '정가 50,000원/명 → 40,000원 (1만원 할인)')
        data_row('쇼업 보장', '리드젠 + 참석 보장', f"게런티 {r['g']}명 × 인당 350,000원",
                 350000, r['g'], '정가 450,000원/명 → 350,000원 (10만원 할인)')
    if r['genCount'] > 0:
        cell_ref['gen'] = data_row('참관객 응대', '초청·리마인드·확정 컨택', f"참관객 {r['genCount']}명 × 인당 20,000원",
                                   20000, r['genCount'], '쇼업 게런티 아님 · 게런티와 독립')
    if r['g'] == 0 and r['genCount'] == 0:
        first = row  # 빈 섹션 — 소계 0
        sec_sub[7] = subtotal(row, row - 1)
    else:
        sec_sub[7] = subtotal(first, row - 1)

    # ---- 푸터
    ws.merge_cells(f'A{row}:H{row}')
    ws[f'A{row}'] = '※ 본 견적은 참고용이며, 베뉴 컨디션에 따라 비용은 변동될 수 있습니다. 정확한 견적을 위해 1~2일 소요됩니다.'
    _apply(ws[f'A{row}'], 'footer')

    # ---- 2-pass: 섹션 헤더 참조·PCO 베이스·상단 합계
    for no, hr in ((1, h1), (2, h2), (3, h3), (4, h4), (5, h5), (6, h6), (7, h7)):
        ws[f'F{hr}'] = f'=F{sec_sub[no]}'
    base_refs = [f'F{sec_sub[i]}' for i in (1, 2, 3, 4, 5)]
    if 'rsvp' in cell_ref:
        base_refs.append(f'F{cell_ref["rsvp"]}')
    if 'gen' in cell_ref:
        base_refs.append(f'F{cell_ref["gen"]}')
    ws[f'D{pco_row}'] = '=' + '+'.join(base_refs)

    total_refs = '+'.join(f'F{sec_sub[i]}' for i in range(1, 8))
    noopt_base = '+'.join(f'F{sec_sub[i]}' for i in (1, 2, 3, 4))
    noopt_pco_base = noopt_base + (f'+F{cell_ref["rsvp"]}' if 'rsvp' in cell_ref else '') \
                                + (f'+F{cell_ref["gen"]}' if 'gen' in cell_ref else '')

    for rr, (label, dfml) in {
        10: ('총 견적금액(VAT별도)', f'={total_refs}'),
        11: ('추가옵션 제외(VAT별도)', f'={noopt_base}+FLOOR(({noopt_pco_base})*0.25,10000)+F{sec_sub[7]}'),
    }.items():
        ws.merge_cells(f'B{rr}:C{rr}')
        ws.merge_cells(f'D{rr}:E{rr}')
        ws.merge_cells(f'G{rr}:H{rr}')
        ws[f'A{rr}'] = label; _apply(ws[f'A{rr}'], 'label_o')
        _apply(ws[f'B{rr}'], 'tot_kor')
        ws[f'D{rr}'] = dfml; _apply(ws[f'D{rr}'], 'tot_val')
        ws[f'F{rr}'] = 'VAT 포함'; _apply(ws[f'F{rr}'], 'tot_sub')
        ws[f'G{rr}'] = f'=ROUND(D{rr}*1.1,0)'; _apply(ws[f'G{rr}'], 'tot_vat')
    ws.row_dimensions[10].height = 27.5
    if use_numberstring:
        ws['B10'] = '="일금 "&NUMBERSTRING(D10,1)&"원 정 ("&TEXT(D10,"#,##0")&"원)"'
        ws['B11'] = '="일금 "&NUMBERSTRING(D11,1)&"원 정 ("&TEXT(D11,"#,##0")&"원)"'
    else:
        ws['B10'] = f'="일금 {_kor(r["pk"])}원 정 ("&TEXT(D10,"#,##0")&"원)"'
        ws['B11'] = f'="일금 {_kor(r["pk_excluding_options"])}원 정 ("&TEXT(D11,"#,##0")&"원)"'

    wb.save(out_path)
    return out_path


def verify(path, result):
    """recalc 이후 호출: D10↔pk, D11↔pk_excluding_options 0원 일치 확인."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    ok = (ws['D10'].value == result['pk']) and (ws['D11'].value == result['pk_excluding_options'])
    return ok, {'D10': ws['D10'].value, 'pk': result['pk'],
                'D11': ws['D11'].value, 'pk_noopt': result['pk_excluding_options']}


# ---------- 재계산 (recalc) ----------
def find_recalc_script():
    """xlsx 스킬 recalc.py 탐색. 순서: 환경변수 XLSX_RECALC(파일 경로) → <skills>/xlsx/scripts/recalc.py(형제)
    → ~/.claude/skills/xlsx/scripts/recalc.py → ~/.claude/skills/synced/*/xlsx/scripts/recalc.py. 없으면 None."""
    cands = []
    env = os.environ.get('XLSX_RECALC')
    if env:
        cands.append(Path(env).expanduser())
    here = Path(__file__).resolve()
    if len(here.parents) >= 3:
        cands.append(here.parents[2] / 'xlsx' / 'scripts' / 'recalc.py')
    home = Path.home()
    cands.append(home / '.claude' / 'skills' / 'xlsx' / 'scripts' / 'recalc.py')
    cands.extend(sorted(home.glob('.claude/skills/synced/*/xlsx/scripts/recalc.py')))
    for c in cands:
        if c.is_file():
            return c
    return None


_MANUAL_HINT = 'Excel에서 파일을 열어 저장(Ctrl+S)한 뒤 verify(path, result)로 D10↔pk 일치를 확인하세요.'


def _recalc_cell_errors(stdout):
    """xlsx 스킬 recalc.py의 JSON 출력에서 셀 오류(#NAME?·#REF! 등) 위치를 뽑는다.
    recalc.py는 셀 오류가 있어도 rc=0·status='errors_found'로 돌려주므로 종료코드만으로는 걸러지지 않는다.
    반환: {오류문자열: [시트!셀, ...]} — 없거나 JSON이 아니면 {}."""
    try:
        data = json.loads((stdout or '').strip())
    except ValueError:
        return {}
    if not isinstance(data, dict) or data.get('status') != 'errors_found':
        return {}
    return {k: list((v or {}).get('locations', [])) for k, v in (data.get('error_summary') or {}).items() if v}


def _soffice_recalc(exe, path, timeout):
    """LibreOffice headless 변환(xlsx→xlsx)으로 수식 캐시값을 채우고 원본을 교체한다. 실패 시 원본 보존, False."""
    src = Path(path).resolve()
    with tempfile.TemporaryDirectory(prefix='estimate-recalc-') as tmp:
        outdir = Path(tmp) / 'out'
        outdir.mkdir()
        profile = Path(tmp) / 'profile'   # 전용 프로필 — 실행 중인 다른 LibreOffice 인스턴스와 충돌 방지
        cmd = [exe, '--headless', '--norestore', f'-env:UserInstallation={profile.as_uri()}',
               '--convert-to', 'xlsx', '--outdir', str(outdir), str(src)]
        try:
            cp = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except (subprocess.TimeoutExpired, OSError) as e:
            print(f'[recalc] LibreOffice 실행 실패({e}) — 원본 유지. {_MANUAL_HINT}')
            return False
        produced = outdir / src.name
        if cp.returncode != 0 or not produced.is_file() or produced.stat().st_size == 0:
            detail = (cp.stderr or cp.stdout or '').strip()[:300]
            print(f'[recalc] LibreOffice 변환 실패(rc={cp.returncode}) — 원본 유지. {detail} {_MANUAL_HINT}')
            return False
        try:
            openpyxl.load_workbook(str(produced), read_only=True).close()   # 깨진 결과로 원본을 덮지 않도록 선검사
        except Exception as e:  # noqa: BLE001
            print(f'[recalc] 변환 결과 열기 실패({e}) — 원본 유지. {_MANUAL_HINT}')
            return False
        staged = src.with_name(src.name + '.recalc~')
        try:
            shutil.copyfile(str(produced), str(staged))   # 같은 폴더에 올린 뒤
            os.replace(str(staged), str(src))             # 원자적 교체 — 중간 실패 시 원본 그대로
        except OSError as e:
            if staged.exists():
                try:
                    staged.unlink()
                except OSError:
                    pass
            print(f'[recalc] 원본 교체 실패({e}) — 원본 유지. {_MANUAL_HINT}')
            return False
    print(f'[recalc] OK — LibreOffice headless: {exe}')
    return True


def recalc(path, timeout=180):
    """수식 재계산 → 성공 시 True.
    ① xlsx 스킬 recalc.py가 있으면 [sys.executable, recalc.py, path] 실행
    ② 없으면(또는 실패하면) soffice/libreoffice --headless --convert-to xlsx 로 변환 후 원본 교체
    ③ 둘 다 없으면 한국어 안내(Excel에서 열어 저장 후 verify) 출력 후 False"""
    path = str(path)
    if not os.path.isfile(path):
        print(f'[recalc] 파일 없음: {path}')
        return False
    script = find_recalc_script()
    script_failed = False
    if script is not None:
        try:
            cp = subprocess.run([sys.executable, str(script), path], capture_output=True, text=True, timeout=timeout)
            if cp.returncode == 0:
                errs = _recalc_cell_errors(cp.stdout)
                if errs:
                    n = sum(len(v) for v in errs.values())
                    where = ' · '.join(f"{k} {', '.join(v[:5])}" for k, v in errs.items())
                    print(f'[recalc] 수식 오류 {n}건: {where} — 재계산은 됐으나 셀 오류가 남아 있습니다. '
                          'Excel에서 확인(Step 6.5 게이트 3). NUMBERSTRING(use_numberstring=True)은 LibreOffice 미지원.')
                    return False
                print(f'[recalc] OK — xlsx 스킬 recalc.py: {script}')
                return True
            script_failed = True
            detail = (cp.stderr or cp.stdout or '').strip()[:300]
            print(f'[recalc] recalc.py 실패(rc={cp.returncode}) → LibreOffice 직접 변환으로 폴백. {detail}')
        except (subprocess.TimeoutExpired, OSError) as e:
            script_failed = True
            print(f'[recalc] recalc.py 실행 오류({e}) → LibreOffice 직접 변환으로 폴백')
    exe = shutil.which('soffice') or shutil.which('libreoffice')
    if exe:
        return _soffice_recalc(exe, path, timeout)
    if script_failed:
        print(f'[recalc] recalc.py 실패 + LibreOffice(soffice/libreoffice) 미탐지 — {_MANUAL_HINT}')
    else:
        print('[recalc] 재계산 도구 없음 — xlsx 스킬 recalc.py(XLSX_RECALC·형제·~/.claude/skills)와 '
              f'LibreOffice(soffice/libreoffice) 모두 미탐지. {_MANUAL_HINT}')
    return False


# ---------- 자가 테스트 ----------
def _self_test():
    """calc(target=100, guarantee=100) → 임시 폴더 export → recalc → verify. recalc 불가 환경이면 verify SKIP.
    종료코드 0=통과 / 1=실패."""
    from calc_estimate_remember import calc_estimate_remember
    print(f"palette : {_P['_source']}")
    r = calc_estimate_remember({'target': 100, 'guarantee': 100})
    with tempfile.TemporaryDirectory(prefix='estimate-selftest-') as tmp:
        out = os.path.join(tmp, '리멤버견적서_selftest.xlsx')
        export_remember_estimate(r, {'project_title': 'self-test'}, out)
        print(f'export  : OK ({os.path.getsize(out):,} bytes) pk={r["pk"]:,} pk_noopt={r["pk_excluding_options"]:,}')
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        nform = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('='))
        d10 = ws['D10'].value
        wb.close()
        if nform == 0 or not (isinstance(d10, str) and d10.startswith('=')):
            print(f'reopen  : FAIL (수식 {nform}개, D10={d10!r})')
            return 1
        print(f'reopen  : OK (수식 {nform}개, D10={d10})')
        if not recalc(out):
            print('recalc  : 불가 → verify SKIP (Excel에서 열어 저장 후 verify)')
            print('export_estimate_remember self-test PASS (verify SKIP)')
            return 0
        ok, info = verify(out, r)
        print(f"verify  : {'OK' if ok else 'FAIL'} {info}")
        print('export_estimate_remember self-test', 'PASS' if ok else 'FAIL')
        return 0 if ok else 1


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(description='리멤버 견적서 xlsx 브리지 — 라이브러리 모듈. CLI는 자가 테스트만 제공.')
    ap.add_argument('--self-test', action='store_true', help='calc → export → recalc → verify 자가 테스트')
    a = ap.parse_args()
    if a.self_test:
        raise SystemExit(_self_test())
    ap.print_help()
