#!/usr/bin/env python3
"""
build_matrix.py - RFP 평가 매트릭스 .xlsx 생성

7개 시트로 구성된 평가 매트릭스를 jc-design-system 토큰을 적용하여 생성한다.

사용:
    from build_matrix import build_evaluation_matrix
    output_path = build_evaluation_matrix(
        analysis_data,
        output_dir="/mnt/user-data/outputs"
    )
"""

import os
import sys
from datetime import datetime
from typing import Dict, List
from pathlib import Path

# Windows 콘솔 한글 출력 안정화 (UTF-8 강제, Sprint 5 BL-S2 패턴)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.formatting.rule import (
    CellIsRule, FormulaRule, ColorScaleRule
)
from openpyxl.utils import get_column_letter


# =====================================================================
# 1. 디자인 토큰 (jc-design-system 호출)
# =====================================================================
HEX_PRIMARY     = "0A2540"
HEX_ACCENT      = "2962FF"
HEX_NEON        = "00E676"
HEX_ORANGE      = "FF5722"
HEX_MAGENTA     = "E91E63"
HEX_LIGHT_NAVY  = "F0F4FA"
HEX_LIGHT_GRAY  = "F8F9FB"
HEX_BORDER      = "E0E0E0"
HEX_DARK_GRAY   = "333333"
HEX_MID_GRAY    = "777777"
HEX_WHITE       = "FFFFFF"

FONT_NAME = "Pretendard"


# =====================================================================
# 2. 스타일 헬퍼
# =====================================================================
def _style_header(cell):
    """헤더 행 표준 스타일"""
    cell.font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_WHITE)
    cell.fill = PatternFill("solid", fgColor=HEX_PRIMARY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color=HEX_WHITE),
        right=Side(style="thin", color=HEX_WHITE),
        top=Side(style="thin", color=HEX_PRIMARY),
        bottom=Side(style="thin", color=HEX_PRIMARY),
    )


def _style_data(cell, alt_row=False):
    """데이터 행 표준 스타일"""
    cell.font = Font(name=FONT_NAME, size=10, color=HEX_DARK_GRAY)
    if alt_row:
        cell.fill = PatternFill("solid", fgColor=HEX_LIGHT_GRAY)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color=HEX_BORDER),
        right=Side(style="thin", color=HEX_BORDER),
        top=Side(style="thin", color=HEX_BORDER),
        bottom=Side(style="thin", color=HEX_BORDER),
    )


def _style_label(cell):
    """라벨(가운데 정렬, 굵게) 셀"""
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=HEX_DARK_GRAY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color=HEX_BORDER),
        right=Side(style="thin", color=HEX_BORDER),
        top=Side(style="thin", color=HEX_BORDER),
        bottom=Side(style="thin", color=HEX_BORDER),
    )


def _set_col_widths(ws, widths: List[float]):
    """컬럼 너비 일괄 설정"""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_headers(ws, row: int, headers: List[str], widths: List[float] = None):
    """헤더 행 작성"""
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        _style_header(c)
    ws.row_dimensions[row].height = 28
    if widths:
        _set_col_widths(ws, widths)


def _write_data_rows(ws, start_row: int, data: List[List]):
    """데이터 행 작성"""
    for r_offset, row_data in enumerate(data):
        row = start_row + r_offset
        ws.row_dimensions[row].height = 22
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row, column=c_idx, value=val)
            _style_data(cell, alt_row=(r_offset % 2 == 0))


# =====================================================================
# 3. 시트 빌더 — 0_종합
# =====================================================================
def _build_summary_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("0_종합", 0)
    _set_col_widths(ws, [25, 40, 30])

    # 타이틀 (병합)
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = "RFP 분석 평가 매트릭스"
    c.font = Font(name=FONT_NAME, size=16, bold=True, color=HEX_WHITE)
    c.fill = PatternFill("solid", fgColor=HEX_PRIMARY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 42

    # 메타 정보
    meta = [
        ("발주처", data.get("client", "-"), ""),
        ("행사명", data.get("event_name", "-"), ""),
        ("분석 일자", datetime.now().strftime("%Y-%m-%d"), ""),
        ("분석자", data.get("analyst", "MICE 전략가"), ""),
    ]
    for i, (label, val, note) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=val)
        ws.cell(row=i, column=3, value=note)
        for col in range(1, 4):
            _style_data(ws.cell(row=i, column=col), alt_row=(i % 2 == 0))
        # 라벨 컬럼은 굵게
        ws.cell(row=i, column=1).font = Font(name=FONT_NAME, size=10, bold=True, color=HEX_DARK_GRAY)

    # 7축 점수
    ws.cell(row=8, column=1, value="▼ 7축 종합 점수").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)

    axes = data.get("axes_scores", {})
    axis_data = [
        ("1. 요건",     axes.get("requirements", 0), "가중치 15%"),
        ("2. 평가",     axes.get("evaluation", 0),   "가중치 20%"),
        ("3. 리스크",   axes.get("risks", 0),        "가중치 20%"),
        ("4. 경쟁",     axes.get("competition", 0),  "가중치 15%"),
        ("5. 일정",     axes.get("timeline", 0),     "가중치 10%"),
        ("6. 예산",     axes.get("budget", 0),       "가중치 15%"),
        ("7. 종합 가치", axes.get("strategy", 0),    "가중치 5%"),
    ]
    for i, (label, score, note) in enumerate(axis_data, start=9):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=score)
        ws.cell(row=i, column=3, value=note)
        for col in range(1, 4):
            _style_data(ws.cell(row=i, column=col), alt_row=(i % 2 == 0))
        ws.cell(row=i, column=1).font = Font(name=FONT_NAME, size=10, bold=True, color=HEX_DARK_GRAY)

    # 가중 평균 종합 점수 (수식)
    ws.cell(row=17, column=1, value="▶ 가중 평균 종합 점수").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)
    ws.cell(row=17, column=2, value="=B9*0.15+B10*0.2+B11*0.2+B12*0.15+B13*0.1+B14*0.15+B15*0.05")
    ws.cell(row=17, column=2).number_format = "0.00"

    # 판정 (수식)
    ws.cell(row=18, column=1, value="▶ 판정").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)
    ws.cell(row=18, column=2, value='=IF(B17>=4,"GO",IF(B17>=3,"GO 조건부",IF(B17>=2,"HOLD","NO-GO")))')
    judge_cell = ws.cell(row=18, column=2)
    judge_cell.font = Font(name=FONT_NAME, size=12, bold=True, color=HEX_WHITE)
    judge_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[18].height = 30

    # 조건부 서식 — 판정 셀
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add(
        "B18",
        FormulaRule(formula=['$B$18="GO"'], fill=PatternFill("solid", fgColor=HEX_NEON))
    )
    ws.conditional_formatting.add(
        "B18",
        FormulaRule(formula=['$B$18="GO 조건부"'], fill=PatternFill("solid", fgColor=HEX_NEON))
    )
    ws.conditional_formatting.add(
        "B18",
        FormulaRule(formula=['$B$18="HOLD"'], fill=PatternFill("solid", fgColor=HEX_ORANGE))
    )
    ws.conditional_formatting.add(
        "B18",
        FormulaRule(formula=['$B$18="NO-GO"'], fill=PatternFill("solid", fgColor=HEX_MAGENTA))
    )

    # 추정 승률
    ws.cell(row=19, column=1, value="▶ 추정 승률").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)
    ws.cell(row=19, column=2, value=data.get("winrate_text", "-"))
    ws.cell(row=19, column=3, value="참고용")

    # 핵심 메시지
    ws.cell(row=21, column=1, value="▼ 핵심 메시지 (GO인 경우)").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)
    messages = data.get("core_messages", [])
    for i, msg in enumerate(messages[:3], start=22):
        ws.cell(row=i, column=1, value=f"메시지 {i-21}")
        ws.cell(row=i, column=2, value=msg)
        for col in range(1, 4):
            _style_data(ws.cell(row=i, column=col), alt_row=(i % 2 == 0))


# =====================================================================
# 4. 시트 빌더 — 1_요건
# =====================================================================
def _build_requirements_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("1_요건")
    headers = ["분류", "요건 항목", "출처(페이지·조항)", "우리 대응", "비고"]
    _write_headers(ws, 1, headers, widths=[12, 40, 18, 18, 30])

    rows = data.get("requirements_rows", [])
    _write_data_rows(ws, 2, rows)

    # 조건부 서식 — 분류 컬럼
    last_row = max(2, 1 + len(rows))
    range_a = f"A2:A{last_row}"
    ws.conditional_formatting.add(
        range_a,
        FormulaRule(formula=['$A2="필수"'],
                    fill=PatternFill("solid", fgColor=HEX_MAGENTA),
                    font=Font(name=FONT_NAME, color=HEX_WHITE, bold=True))
    )
    ws.conditional_formatting.add(
        range_a,
        FormulaRule(formula=['$A2="선택"'],
                    fill=PatternFill("solid", fgColor=HEX_ORANGE))
    )
    ws.conditional_formatting.add(
        range_a,
        FormulaRule(formula=['$A2="가산"'],
                    fill=PatternFill("solid", fgColor=HEX_NEON))
    )


# =====================================================================
# 5. 시트 빌더 — 2_평가
# =====================================================================
def _build_evaluation_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("2_평가")
    headers = ["대분류", "평가 항목", "배점", "가중치(%)", "예상 점수", "환산 점수", "강점 여부", "비고"]
    _write_headers(ws, 1, headers, widths=[18, 35, 10, 12, 14, 14, 14, 25])

    rows = data.get("evaluation_rows", [])
    _write_data_rows(ws, 2, rows)

    # G열 강점 조건부 서식
    last_row = max(2, 1 + len(rows))
    range_g = f"G2:G{last_row}"
    ws.conditional_formatting.add(
        range_g,
        FormulaRule(formula=['$G2="강점"'],
                    fill=PatternFill("solid", fgColor=HEX_ACCENT),
                    font=Font(name=FONT_NAME, color=HEX_WHITE, bold=True))
    )
    ws.conditional_formatting.add(
        range_g,
        FormulaRule(formula=['$G2="약점"'],
                    fill=PatternFill("solid", fgColor=HEX_MAGENTA),
                    font=Font(name=FONT_NAME, color=HEX_WHITE, bold=True))
    )


# =====================================================================
# 6. 시트 빌더 — 3_리스크
# =====================================================================
def _build_risks_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("3_리스크")
    headers = ["영역", "조항 원문", "출처", "등급", "영향 분석", "대응 방안"]
    _write_headers(ws, 1, headers, widths=[15, 50, 18, 12, 30, 30])

    rows = data.get("risk_rows", [])
    _write_data_rows(ws, 2, rows)

    # D열 등급 조건부 서식
    last_row = max(2, 1 + len(rows))
    range_d = f"D2:D{last_row}"
    ws.conditional_formatting.add(
        range_d,
        FormulaRule(formula=['$D2="상"'],
                    fill=PatternFill("solid", fgColor=HEX_MAGENTA),
                    font=Font(name=FONT_NAME, color=HEX_WHITE, bold=True))
    )
    ws.conditional_formatting.add(
        range_d,
        FormulaRule(formula=['$D2="중"'],
                    fill=PatternFill("solid", fgColor=HEX_ORANGE))
    )
    ws.conditional_formatting.add(
        range_d,
        FormulaRule(formula=['$D2="하"'],
                    fill=PatternFill("solid", fgColor=HEX_NEON))
    )


# =====================================================================
# 7. 시트 빌더 — 4_경쟁
# =====================================================================
def _build_competition_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("4_경쟁")
    headers = ["예상 경쟁사", "강점(예상)", "약점(예상)", "위협도", "우리 대비"]
    _write_headers(ws, 1, headers, widths=[25, 35, 35, 14, 14])

    rows = data.get("competition_rows", [])
    _write_data_rows(ws, 2, rows)

    # SWOT 영역 (시트 하단)
    swot_start = max(8, len(rows) + 4)
    ws.cell(row=swot_start, column=1, value="▼ 우리 측 SWOT").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)

    swot = data.get("swot", {})
    swot_items = [
        ("강점", swot.get("strength", "-")),
        ("약점", swot.get("weakness", "-")),
        ("기회", swot.get("opportunity", "-")),
        ("위협", swot.get("threat", "-")),
    ]
    for i, (label, val) in enumerate(swot_items, start=swot_start + 1):
        ws.cell(row=i, column=1, value=label)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
        ws.cell(row=i, column=2, value=val)
        for col in range(1, 6):
            _style_data(ws.cell(row=i, column=col), alt_row=(i % 2 == 0))


# =====================================================================
# 8. 시트 빌더 — 5_일정
# =====================================================================
def _build_timeline_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("5_일정")
    headers = ["단계", "일자", "D-Day", "직전 간격(일)", "압박 강도", "비고"]
    _write_headers(ws, 1, headers, widths=[18, 14, 10, 18, 14, 25])

    rows = data.get("timeline_rows", [])
    _write_data_rows(ws, 2, rows)

    # E열 압박 강도 조건부 서식
    last_row = max(2, 1 + len(rows))
    range_e = f"E2:E{last_row}"
    ws.conditional_formatting.add(
        range_e,
        FormulaRule(formula=['$E2="상"'],
                    fill=PatternFill("solid", fgColor=HEX_MAGENTA),
                    font=Font(name=FONT_NAME, color=HEX_WHITE, bold=True))
    )
    ws.conditional_formatting.add(
        range_e,
        FormulaRule(formula=['$E2="중"'],
                    fill=PatternFill("solid", fgColor=HEX_ORANGE))
    )
    ws.conditional_formatting.add(
        range_e,
        FormulaRule(formula=['$E2="하"'],
                    fill=PatternFill("solid", fgColor=HEX_NEON))
    )


# =====================================================================
# 9. 시트 빌더 — 6_예산
# =====================================================================
def _build_budget_sheet(wb: Workbook, data: Dict):
    ws = wb.create_sheet("6_예산")
    headers = ["항목", "금액(원)", "비고"]
    _write_headers(ws, 1, headers, widths=[25, 18, 35])

    budget = data.get("budget_data", {})
    rows = [
        ("발주가 (총사업비)",     budget.get("announced", 0),         ""),
        ("부가세",               budget.get("vat", "-"),              "별도/포함"),
        ("추정 원가",             budget.get("estimated_cost", 0),    "자체 산정"),
        ("추정 마진",             "=B2-B4",                           "발주가 - 원가"),
        ("추정 마진율",           "=B5/B2",                           "표준 15~25%"),
        ("MICE 표준 마진율",      budget.get("standard_margin", 0.18), "참조 값"),
        ("표준 대비 차이",        "=B6-B7",                           ""),
    ]
    for i, (label, val, note) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=val)
        ws.cell(row=i, column=3, value=note)
        for col in range(1, 4):
            _style_data(ws.cell(row=i, column=col), alt_row=(i % 2 == 0))
        ws.cell(row=i, column=1).font = Font(name=FONT_NAME, size=10, bold=True, color=HEX_DARK_GRAY)
    ws.cell(row=2, column=2).number_format = "#,##0"
    ws.cell(row=4, column=2).number_format = "#,##0"
    ws.cell(row=5, column=2).number_format = "#,##0"
    ws.cell(row=6, column=2).number_format = "0.00%"
    ws.cell(row=7, column=2).number_format = "0.00%"
    ws.cell(row=8, column=2).number_format = "0.00%"

    # 입찰가 시뮬레이션
    ws.cell(row=10, column=1, value="▼ 입찰가 권고 시뮬레이션").font = Font(name=FONT_NAME, size=11, bold=True, color=HEX_PRIMARY)
    sim_rows = [
        ("목표 마진율", budget.get("target_margin", 0.18), "행사 유형별 표준"),
        ("리스크 프리미엄", budget.get("risk_premium", 0.05), "리스크 시트 참조"),
        ("입찰가 권고", "=B4*(1+B11)*(1+B12)", "원가 × (1+마진) × (1+프리미엄)"),
        ("발주가 대비", "=B13/B2", ""),
    ]
    for i, (label, val, note) in enumerate(sim_rows, start=11):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=val)
        ws.cell(row=i, column=3, value=note)
        for col in range(1, 4):
            _style_data(ws.cell(row=i, column=col), alt_row=(i % 2 == 0))
        ws.cell(row=i, column=1).font = Font(name=FONT_NAME, size=10, bold=True, color=HEX_DARK_GRAY)
    ws.cell(row=11, column=2).number_format = "0.00%"
    ws.cell(row=12, column=2).number_format = "0.00%"
    ws.cell(row=13, column=2).number_format = "#,##0"
    ws.cell(row=14, column=2).number_format = "0.00%"

    # B6 마진율 조건부 서식
    ws.conditional_formatting.add(
        "B6",
        CellIsRule(operator="lessThan", formula=["0.05"],
                   fill=PatternFill("solid", fgColor=HEX_MAGENTA))
    )
    ws.conditional_formatting.add(
        "B6",
        CellIsRule(operator="between", formula=["0.05", "0.15"],
                   fill=PatternFill("solid", fgColor=HEX_ORANGE))
    )
    ws.conditional_formatting.add(
        "B6",
        CellIsRule(operator="greaterThanOrEqual", formula=["0.15"],
                   fill=PatternFill("solid", fgColor=HEX_NEON))
    )


# =====================================================================
# 10. 메인 빌더
# =====================================================================
def build_evaluation_matrix(
    data: Dict,
    output_dir: str = "/mnt/user-data/outputs"
) -> str:
    """
    7축 분석 데이터를 .xlsx 파일로 빌드.
    
    Args:
        data: 분석 결과 dict
        output_dir: 출력 디렉토리
    
    Returns:
        생성된 파일 경로
    """
    wb = Workbook()
    # 기본 시트 제거
    wb.remove(wb.active)

    # 7개 시트 빌드 (순서 유지)
    _build_summary_sheet(wb, data)
    _build_requirements_sheet(wb, data)
    _build_evaluation_sheet(wb, data)
    _build_risks_sheet(wb, data)
    _build_competition_sheet(wb, data)
    _build_timeline_sheet(wb, data)
    _build_budget_sheet(wb, data)

    # 파일 저장
    today = datetime.now().strftime("%Y%m%d")
    client = data.get("client", "client").replace(" ", "_")
    event = data.get("event_name", "event").replace(" ", "_")
    filename = f"rfp-evaluation-matrix_{client}_{event}_{today}.xlsx"
    output_path = os.path.join(output_dir, filename)
    os.makedirs(output_dir, exist_ok=True)
    wb.save(output_path)
    return output_path


# =====================================================================
# 11. CLI 진입점 (검증용)
# =====================================================================
if __name__ == "__main__":
    sample = {
        "client": "샘플발주처",
        "event_name": "샘플 컨퍼런스",
        "analyst": "MICE 전략가",
        "axes_scores": {
            "requirements": 4.5, "evaluation": 4.2, "risks": 4.0,
            "competition": 4.0, "timeline": 4.5, "budget": 4.0, "strategy": 4.5
        },
        "winrate_text": "45% (참고용)",
        "core_messages": [
            "18년 경력 검증된 운영 안정성",
            "동일 발주처 유사 행사 실적",
            "[파트너 협업] 차별화"
        ],
        "requirements_rows": [
            ["필수", "PCO 자격 보유", "p.3 2조", "가능", ""],
            ["필수", "유사 실적 3건 이상", "p.3 2조", "가능", "포트폴리오 4건 보유"],
            ["선택", "외국어 운영 가능", "p.5 4조", "조건부", "통역사 외주"],
            ["가산", "지역 업체 가산점", "p.7 6조", "불가", "본사 서울"],
        ],
        "evaluation_rows": [
            ["기술", "운영 계획", 30, 30, 28, 28, "강점", ""],
            ["기술", "차별화 방안", 20, 20, 18, 18, "강점", ""],
            ["기술", "조직 구성", 15, 15, 13, 13, "중립", ""],
            ["기술", "유사 실적", 15, 15, 14, 14, "강점", ""],
            ["가격", "입찰 가격", 20, 20, 16, 16, "중립", ""],
        ],
        "risk_rows": [
            ["손해배상", "직접·간접 손해 배상", "p.9 12조", "중", "간접 손해 무한 확장 가능", "한도 명시 협상"],
            ["변경권", "발주처 자유 변경권", "p.10 14조", "중", "추가 비용 보상 부재", "별도 협의 조항 추가"],
        ],
        "competition_rows": [
            ["A사", "발주처 친밀도", "단가 경쟁력 부족", "중", "동등"],
            ["B사", "기술력", "유사 실적 부족", "중", "우위"],
        ],
        "swot": {
            "strength": "18년 경력 + [파트너 협업]",
            "weakness": "지역 가산점 미충족",
            "opportunity": "발주처 신규 사업 확장 의향",
            "threat": "A사 친밀도"
        },
        "timeline_rows": [
            ["공고", "2026-05-01", 0, 0, "하", ""],
            ["질의 마감", "2026-05-15", 0, 14, "하", ""],
            ["제안 마감", "2026-06-01", 0, 17, "하", ""],
            ["발표", "2026-06-15", 0, 14, "하", ""],
            ["계약", "2026-07-01", 0, 16, "하", ""],
            ["행사일", "2026-09-15", 0, 76, "하", ""],
        ],
        "budget_data": {
            "announced": 300_000_000,
            "vat": "별도",
            "estimated_cost": 246_000_000,
            "standard_margin": 0.18,
            "target_margin": 0.18,
            "risk_premium": 0.05
        }
    }
    path = build_evaluation_matrix(sample, output_dir="/tmp")
    print(f"생성 완료: {path}")
