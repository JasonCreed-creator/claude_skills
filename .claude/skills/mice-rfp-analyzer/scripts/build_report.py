#!/usr/bin/env python3
"""
build_report.py - RFP 분석 보고서 .docx 생성

mice-rfp-analyzer 의 7축 분석 결과를 받아 jc-design-system 토큰이 적용된
.docx 분석 보고서를 생성한다.

사용:
    from build_report import build_analysis_report
    output_path = build_analysis_report(
        analysis_data,
        output_dir="/mnt/user-data/outputs",
        mode="full"  # or "quick"
    )
"""

import os
import sys
from datetime import datetime
from typing import Dict
from pathlib import Path

# Windows 콘솔 한글 출력 안정화 (UTF-8 강제, Sprint 5 BL-S2 패턴)
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Mm, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# =====================================================================
# 1. 디자인 토큰 (jc-design-system 호출)
# =====================================================================
COLOR_PRIMARY      = RGBColor(0x0A, 0x25, 0x40)  # Deep Navy
COLOR_ACCENT       = RGBColor(0x29, 0x62, 0xFF)  # Electric Blue
COLOR_NEON         = RGBColor(0x00, 0xE6, 0x76)  # Neon Green (GO)
COLOR_ORANGE       = RGBColor(0xFF, 0x57, 0x22)  # Orange (HOLD)
COLOR_MAGENTA      = RGBColor(0xE9, 0x1E, 0x63)  # Magenta (NO-GO)
COLOR_DARK_GRAY    = RGBColor(0x33, 0x33, 0x33)
COLOR_MID_GRAY     = RGBColor(0x77, 0x77, 0x77)
COLOR_WHITE        = RGBColor(0xFF, 0xFF, 0xFF)

# Hex strings (셀 배경용)
HEX_PRIMARY     = "0A2540"
HEX_ACCENT      = "2962FF"
HEX_NEON        = "00E676"
HEX_ORANGE      = "FF5722"
HEX_MAGENTA     = "E91E63"
HEX_LIGHT_NAVY  = "F0F4FA"
HEX_LIGHT_GRAY  = "F8F9FB"
HEX_BORDER      = "E0E0E0"

FONT_KO = "Pretendard"
FONT_FALLBACK = "맑은 고딕"


# =====================================================================
# 2. 유틸리티 함수
# =====================================================================
def _set_cell_bg(cell, hex_color: str):
    """셀 배경색 설정"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_run_font(run, name=FONT_KO, size=10.5, bold=False, color=COLOR_DARK_GRAY):
    """런 글꼴 일괄 설정"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # 한글 글꼴 매핑
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def _set_para_spacing(para, before=6, after=6, line=1.4):
    """단락 간격 설정"""
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _add_section_break(doc):
    """섹션 구분 (페이지 나눔)"""
    doc.add_page_break()


def _add_heading(doc, text: str, level: int = 1):
    """
    제목 위계 추가
    level 1: 18pt ExtraBold Primary Navy
    level 2: 14pt SemiBold Primary Navy
    level 3: 12pt SemiBold Dark Gray
    """
    para = doc.add_paragraph()
    run = para.add_run(text)
    if level == 1:
        _set_run_font(run, size=18, bold=True, color=COLOR_PRIMARY)
        _set_para_spacing(para, before=24, after=12)
    elif level == 2:
        _set_run_font(run, size=14, bold=True, color=COLOR_PRIMARY)
        _set_para_spacing(para, before=18, after=9)
    else:
        _set_run_font(run, size=12, bold=True, color=COLOR_DARK_GRAY)
        _set_para_spacing(para, before=12, after=6)
    return para


def _add_body(doc, text: str, bold=False, color=COLOR_DARK_GRAY):
    """본문 단락 추가"""
    para = doc.add_paragraph()
    run = para.add_run(text)
    _set_run_font(run, size=10.5, bold=bold, color=color)
    _set_para_spacing(para, before=4, after=4, line=1.4)
    return para


def _add_emphasis_box(doc, text: str):
    """강조 박스 (좌측 4pt Primary Navy 테두리, Light Navy 배경)"""
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _set_cell_bg(cell, HEX_LIGHT_NAVY)

    # 좌측 테두리 4pt Primary Navy
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "32")  # 4pt = 32 1/8th-pt units
    left.set(qn("w:color"), HEX_PRIMARY)
    tcBorders.append(left)
    tcPr.append(tcBorders)

    para = cell.paragraphs[0]
    run = para.add_run(text)
    _set_run_font(run, size=10.5, color=COLOR_DARK_GRAY)
    _set_para_spacing(para, before=8, after=8, line=1.4)


def _add_judgment_badge(doc, judgment: str, score: float):
    """GO/HOLD/NO-GO 색상 박스"""
    color_map = {
        "GO":         (HEX_NEON, COLOR_WHITE),
        "GO 조건부":   (HEX_NEON, COLOR_PRIMARY),
        "HOLD":       (HEX_ORANGE, COLOR_WHITE),
        "NO-GO":      (HEX_MAGENTA, COLOR_WHITE),
    }
    bg, fg = color_map.get(judgment, (HEX_PRIMARY, COLOR_WHITE))

    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _set_cell_bg(cell, bg)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f"판정: {judgment}    |    종합 점수: {score:.1f} / 5.0")
    _set_run_font(run, size=14, bold=True, color=fg)
    _set_para_spacing(para, before=12, after=12)


def _add_data_table(doc, headers: list, rows: list, col_widths_cm=None):
    """
    표준 데이터 표 생성
    헤더: Primary Navy 배경 + 흰 글자
    데이터 행: 짝수 Light Gray, 홀수 흰색
    """
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 헤더 행
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        _set_cell_bg(cell, HEX_PRIMARY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(h)
        _set_run_font(run, size=10, bold=True, color=COLOR_WHITE)
        _set_para_spacing(para, before=4, after=4)

    # 데이터 행
    for r_idx, row in enumerate(rows):
        tr = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row):
            cell = tr.cells[c_idx]
            if r_idx % 2 == 0:
                _set_cell_bg(cell, HEX_LIGHT_GRAY)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            run = para.add_run(str(val))
            _set_run_font(run, size=9.5, color=COLOR_DARK_GRAY)
            _set_para_spacing(para, before=2, after=2)

    # 컬럼 너비 설정
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    return table


# =====================================================================
# 3. 보고서 본체 빌더
# =====================================================================
def _set_doc_margins(doc):
    """A4 + 25mm 여백"""
    section = doc.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)


def _build_cover(doc, data: Dict):
    """표지 페이지"""
    # 상단 여백
    for _ in range(3):
        doc.add_paragraph()

    # 메인 타이틀
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RFP 분석 보고서")
    _set_run_font(run, size=32, bold=True, color=COLOR_PRIMARY)
    _set_para_spacing(p, before=0, after=24)

    # 행사명 + 발주처
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data.get("event_name", "[행사명]"))
    _set_run_font(run, size=18, bold=True, color=COLOR_DARK_GRAY)
    _set_para_spacing(p, before=0, after=6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data.get("client", "[발주처]"))
    _set_run_font(run, size=14, color=COLOR_MID_GRAY)
    _set_para_spacing(p, before=0, after=36)

    # 판정 배지
    _add_judgment_badge(doc, data.get("judgment", "HOLD"), data.get("total_score", 0.0))

    # 메타 정보
    doc.add_paragraph()
    meta_rows = [
        ["분석 일자", datetime.now().strftime("%Y-%m-%d")],
        ["분석자", data.get("analyst", "MICE 전략가")],
        ["추정 승률", f"{data.get('winrate', 0)*100:.0f}% (참고용)"],
        ["행사 일자", data.get("event_date", "-")],
        ["발주가", data.get("budget", "-")],
        ["제안 마감", data.get("submission_deadline", "-")],
    ]
    _add_data_table(doc, ["항목", "내용"], meta_rows, col_widths_cm=[5, 12])


def _build_executive_summary(doc, data: Dict):
    _add_heading(doc, "1. Executive Summary", level=1)

    # 결론 박스
    _add_emphasis_box(doc, f"■ 결론\n{data.get('conclusion', '[판정 사유]')}")
    doc.add_paragraph()

    # 핵심 메시지
    _add_heading(doc, "1-1. 핵심 메시지", level=2)
    messages = data.get("core_messages", [])
    for i, msg in enumerate(messages, start=1):
        _add_body(doc, f"{i}. {msg}")

    # 결정적 변수
    _add_heading(doc, "1-2. 결정적 변수 Top 3", level=2)
    decisives = data.get("decisive_factors", [])
    for i, item in enumerate(decisives, start=1):
        _add_body(doc, f"{i}. {item}")

    # 즉시 대응 사항
    _add_heading(doc, "1-3. 즉시 대응 필요 사항", level=2)
    actions = data.get("immediate_actions", [])
    for item in actions:
        _add_body(doc, f"- {item}")


def _build_rfp_overview(doc, data: Dict):
    _add_heading(doc, "2. RFP 개요", level=1)
    rows = [
        ["발주처", data.get("client", "-")],
        ["행사명", data.get("event_name", "-")],
        ["행사 일자", data.get("event_date", "-")],
        ["행사 장소", data.get("event_venue", "-")],
        ["행사 규모", data.get("event_scale", "-")],
        ["발주가", data.get("budget", "-")],
        ["제안 마감", data.get("submission_deadline", "-")],
        ["응찰 자격", data.get("eligibility", "-")],
    ]
    _add_data_table(doc, ["항목", "내용"], rows, col_widths_cm=[5, 12])


def _build_axis_section(doc, axis_num: int, axis_title: str, section_data: Dict):
    """3-1 ~ 3-7 축별 섹션 빌더 (공통 구조)"""
    _add_heading(doc, f"3-{axis_num}. {axis_title}", level=2)
    if section_data.get("summary"):
        _add_body(doc, section_data["summary"])

    # 표 형식 데이터가 있으면 추가
    if section_data.get("table"):
        headers = section_data["table"]["headers"]
        rows = section_data["table"]["rows"]
        _add_data_table(doc, headers, rows)
        doc.add_paragraph()

    # 추가 코멘트
    if section_data.get("comments"):
        for c in section_data["comments"]:
            _add_body(doc, f"• {c}")


def _build_seven_axis(doc, data: Dict):
    _add_heading(doc, "3. 7축 분석", level=1)

    axes = data.get("axes", {})
    titles = [
        "요건 분석",
        "평가 기준 분석",
        "리스크 분석",
        "경쟁 환경 분석",
        "일정 압박 분석",
        "예산 분석",
        "전략 권고",
    ]
    keys = [
        "requirements", "evaluation", "risks",
        "competition", "timeline", "budget", "strategy"
    ]
    for i, (title, key) in enumerate(zip(titles, keys), start=1):
        _build_axis_section(doc, i, title, axes.get(key, {}))


def _build_appendix(doc, data: Dict):
    _add_heading(doc, "4. 부록", level=1)

    _add_heading(doc, "4-1. 발주처 추가 질의 사항", level=2)
    for q in data.get("queries_to_client", []):
        _add_body(doc, f"- {q}")

    _add_heading(doc, "4-2. 후속 액션 체크리스트", level=2)
    for a in data.get("followup_checklist", []):
        _add_body(doc, f"☐ {a}")


# =====================================================================
# 4. 메인 빌더
# =====================================================================
def build_analysis_report(
    data: Dict,
    output_dir: str = "/mnt/user-data/outputs",
    mode: str = "full"
) -> str:
    """
    7축 분석 데이터를 .docx 파일로 빌드.
    
    Args:
        data: 분석 결과 dict (구조는 chaining-guide.md 참조)
        output_dir: 출력 디렉토리
        mode: "full" or "quick"
    
    Returns:
        생성된 파일 경로
    """
    doc = Document()
    _set_doc_margins(doc)

    # 1. 표지
    _build_cover(doc, data)
    _add_section_break(doc)

    # 2. Executive Summary
    _build_executive_summary(doc, data)
    _add_section_break(doc)

    if mode == "full":
        # 3. RFP 개요
        _build_rfp_overview(doc, data)
        _add_section_break(doc)

        # 4. 7축 분석
        _build_seven_axis(doc, data)
        _add_section_break(doc)

        # 5. 부록
        _build_appendix(doc, data)
    else:
        # Quick 모드: 핵심 3축만
        _add_heading(doc, "3. 핵심 분석 (Quick 모드)", level=1)
        axes = data.get("axes", {})
        _build_axis_section(doc, 1, "요건 분석", axes.get("requirements", {}))
        _build_axis_section(doc, 2, "평가 기준 분석", axes.get("evaluation", {}))
        _build_axis_section(doc, 7, "전략 권고", axes.get("strategy", {}))

    # 파일 저장
    today = datetime.now().strftime("%Y%m%d")
    client = data.get("client", "client").replace(" ", "_")
    event = data.get("event_name", "event").replace(" ", "_")
    filename = f"rfp-analysis-report_{client}_{event}_{today}.docx"
    output_path = os.path.join(output_dir, filename)
    os.makedirs(output_dir, exist_ok=True)
    doc.save(output_path)
    return output_path


# =====================================================================
# 5. CLI 진입점 (단독 실행 / 검증용)
# =====================================================================
if __name__ == "__main__":
    # 검증용 샘플 데이터
    sample = {
        "client": "샘플발주처",
        "event_name": "샘플 컨퍼런스",
        "event_date": "2026-09-15",
        "event_venue": "코엑스",
        "event_scale": "500명",
        "budget": "300,000,000원",
        "submission_deadline": "2026-06-01",
        "eligibility": "최근 3년 내 유사 실적 3건 이상",
        "judgment": "GO",
        "total_score": 4.2,
        "winrate": 0.45,
        "analyst": "MICE 전략가",
        "conclusion": "필수 요건 100% 충족 + 평가 가중치 70%가 우리 강점 영역. 응찰 권장.",
        "core_messages": [
            "18년 경력 검증된 운영 안정성",
            "동일 발주처 유사 행사 수행 실적",
            "[파트너 협업] 통한 차별화된 참가자 모집"
        ],
        "decisive_factors": [
            "평가 가중치 60%가 운영 역량 영역 (우리 강점)",
            "예상 경쟁사 2~3개 (진입 장벽 높음)",
            "추정 마진율 18% (표준 상단)"
        ],
        "immediate_actions": [
            "발주처 질의 마감 전 가산점 항목 확인 필요",
            "[파트너 협업] 데이터 정리 (5월 25일까지)",
            "킥오프 미팅 일정 조율"
        ],
        "axes": {
            "requirements": {
                "summary": "필수 요건 12건, 우리 측 100% 대응 가능. 가산 요건 4건 중 3건 충족 가능.",
                "comments": [
                    "별첨3 운영 매뉴얼 요구사항 정독 필요",
                    "산출물에 사후 보고서 포함 (정산 조건)"
                ]
            },
            "evaluation": {
                "summary": "기술 80 : 가격 20. 정량 평가 60% / 정성 평가 40%.",
                "comments": [
                    "운영 계획 배점 30점 (최대) — 우리 강점 영역",
                    "차별화 가산점 5점 — [파트너 협업] 활용 가능"
                ]
            },
            "risks": {
                "summary": "독소 조항 상 등급 0건, 중 등급 2건 식별. 리스크 프리미엄 5% 권고.",
                "comments": [
                    "[중] 손해배상 한도 명시 없음 → 협상 시 한도 설정 요구",
                    "[중] 변경권 조항 → '추가 비용 별도 협의' 추가 요구"
                ]
            },
            "competition": {
                "summary": "예상 경쟁사 2~3개 추정. 진입 장벽 (3년 실적 3건)이 풀을 좁힘.",
                "comments": ["우리 강점: 동일 발주처 유사 실적, [파트너 협업]"]
            },
            "timeline": {
                "summary": "공고 → 제안 마감 21일 (정상). 계약 → 행사일 90일 이상 (정상).",
                "comments": ["압박 강도 전 단계 '하' — 일정 리스크 낮음"]
            },
            "budget": {
                "summary": "발주가 3억 / 추정 원가 2.46억 / 추정 마진율 18% (표준 상단).",
                "comments": ["입찰가 권고: 2.95억 (리스크 프리미엄 5% 반영)"]
            },
            "strategy": {
                "summary": "종합 점수 4.2 / 5.0. GO 판정. 응찰 권장.",
                "comments": [
                    "자원 투입 강도: 강 (시니어 PM + 디자이너 풀가동)",
                    "발표 자료에 운영 안정성·차별화 메시지 집중"
                ]
            }
        },
        "queries_to_client": [
            "별첨3의 운영 매뉴얼 작성 범위 확인",
            "가산점 항목 평가 방식 (절대/상대) 확인",
            "발표 시 발주처 측 참석자 구성 확인"
        ],
        "followup_checklist": [
            "킥오프 미팅 일정 확정 (5월 12일까지)",
            "[파트너 협업] 데이터 정리 (5월 25일까지)",
            "제안서 초안 완성 (5월 28일까지)",
            "내부 모의 발표 (5월 30일)"
        ]
    }
    path = build_analysis_report(sample, output_dir="/tmp", mode="full")
    print(f"생성 완료: {path}")
