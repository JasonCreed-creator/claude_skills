"""build_runsheet.py — 행사 운영 큐시트(run of show) XLSX 생성기.

입력(event 메타 + cues 리스트) → 10컬럼 큐시트 + 변경이력 시트 xlsx.
mice-estimate xlsx 패턴 재사용(openpyxl + jc-design SoT 런타임 로드 + 무결성 검증).

진입점:
    build_runsheet(event, cues, output_path, changelog=None, verify=True) -> str

핵심 보증:
1. 클록 자동 산출 (start_time + 누적 소요)
2. 시간 무결성 검증 (Σ duration_min == end−start, end_time 제공 시)
3. 10컬럼 그리드 + 변경이력 시트
4. jc-design SoT 토큰 적용 (실패 시 미러 폴백)
5. 저장 후 재오픈 무결성

자가검증:  python build_runsheet.py   (반일 컨퍼런스 self-test)
"""

import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# jc-design-system SoT 토큰 (런타임 로드 + 미러 폴백)
#   값 정본: jc-design-system/references/signature-tokens.md §6 JSON.
#   xlsx는 CSS 변수 불가 → hex 리터럴 불가피. 각 값은 SoT 미러이며 하드코딩 아님.
#   openpyxl 은 '#' 없는 RRGGBB 사용.
# ============================================================
def _load_jc_tokens():
    import json
    import re
    try:
        sot = (Path(__file__).resolve().parents[2]
               / "jc-design-system" / "references" / "signature-tokens.md")
        m = re.search(r"```json\s*\n(.*?)\n```", sot.read_text(encoding="utf-8"), re.S)
        return json.loads(m.group(1)) if m else {}
    except Exception:
        return {}


_JC = _load_jc_tokens()


def _jc(key, fallback):
    c = _JC.get("color", {})
    v = c.get(key) or c.get("point", {}).get(key) or c.get("semantic", {}).get(key) or fallback
    return str(v).lstrip("#")


COLOR_PRIMARY = _jc("primary", "0A2540")        # 헤더·타이틀
COLOR_ACCENT = _jc("accent", "2962FF")          # 컬럼 헤더
COLOR_BORDER = _jc("borderStrong", "C9CFD8")    # 구분면
COLOR_ORANGE = _jc("orange", "FF5722")          # 연출 cue 강조 (point)
COLOR_SURFACE_ALT = _jc("surfaceAlt", "F4F6FA") # 휴식·전환 행
COLOR_WHITE = "FFFFFF"
COLOR_TEXT = _jc("text", "1A1D24")

FONT_TITLE = Font(name="Pretendard", size=14, bold=True, color=COLOR_WHITE)
FONT_META = Font(name="Pretendard", size=10, color=COLOR_WHITE)
FONT_HDR = Font(name="Pretendard", size=11, bold=True, color=COLOR_WHITE)
FONT_BODY = Font(name="Pretendard", size=10, color=COLOR_TEXT)
FONT_BODY_BOLD = Font(name="Pretendard", size=10, bold=True, color=COLOR_TEXT)
FONT_CUE = Font(name="Pretendard", size=10, bold=True, color=COLOR_ORANGE)

FILL_TITLE = PatternFill("solid", fgColor=COLOR_PRIMARY)
FILL_HDR = PatternFill("solid", fgColor=COLOR_ACCENT)
FILL_BREAK = PatternFill("solid", fgColor=COLOR_SURFACE_ALT)

_side = Side(style="thin", color=COLOR_BORDER)
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)

AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 10컬럼 정의 (순서·라벨·너비) — cuesheet-columns.md 정본
COLUMNS = [
    ("Cue#", 6), ("시간", 8), ("세그먼트", 16), ("무대·발표", 18),
    ("Audio", 12), ("Video", 14), ("Light", 14), ("연출 cue", 24),
    ("Owner", 9), ("비고", 18),
]
# cue dict 키 ↔ 컬럼(3번째부터)
CUE_KEYS = ["segment", "stage", "audio", "video", "light", "cue", "owner", "note"]

_BREAK_SEGMENTS = ("휴식", "전환", "셋업", "브레이크")


# ============================================================
# 시간 유틸
# ============================================================
def _to_min(hhmm: str) -> int:
    h, m = str(hhmm).strip().split(":")
    return int(h) * 60 + int(m)


def _to_clock(total: int) -> str:
    """0:00 기준 분 → HH:MM. 24시 초과는 +Nd HH:MM."""
    day, rem = divmod(total, 24 * 60)
    s = f"{rem // 60:02d}:{rem % 60:02d}"
    return s if day == 0 else f"+{day}d {s}"


# ============================================================
# 검증 + 클록 산출
# ============================================================
def _validate_and_clock(event: dict, cues: list) -> list:
    """cue 유효성 검증 + 클록 리스트 반환. 시간 무결성 실패 시 ValueError."""
    if not cues:
        raise ValueError("cues 비어있음")
    if not event.get("start_time"):
        raise ValueError("event.start_time 필수")

    start = _to_min(event["start_time"])
    clocks, t = [], start
    for i, c in enumerate(cues):
        if not c.get("segment"):
            raise ValueError(f"cue[{i}] 'segment' 누락")
        d = c.get("duration_min")
        if not isinstance(d, int) or isinstance(d, bool) or d <= 0:
            raise ValueError(f"cue[{i}]({c.get('segment')}) duration_min 양의 정수 아님: {d!r}")
        clocks.append(_to_clock(t))
        t += d

    total = t - start  # = Σ duration_min
    if event.get("end_time"):
        end_decl = _to_min(event["end_time"])
        if end_decl < start:          # 자정 교차
            end_decl += 24 * 60
        if t != end_decl:
            raise ValueError(
                "시간 무결성 실패 (Critical): "
                f"소요 합 {total}분(종료 {_to_clock(t)}) ≠ 선언 총시간 "
                f"{end_decl - start}분(종료 {event['end_time']}). 차이 {t - end_decl:+d}분."
            )
    return clocks


# ============================================================
# 메인 진입점
# ============================================================
def build_runsheet(event: dict, cues: list, output_path: str,
                   changelog: list = None, verify: bool = True) -> str:
    """큐시트 xlsx 생성.

    Args:
        event: {title, date, venue, start_time, end_time?, version?, client_id?} — 외부 주입
        cues:  [{segment, duration_min, stage, audio, video, light, cue, owner, note}, ...]
        output_path: 출력 xlsx 절대경로
        changelog: [(버전, 일자, 변경), ...]. None이면 [(v{version}, today, '초안 생성')]
        verify: 저장 후 재오픈 무결성 검증
    Returns: output_path
    """
    clocks = _validate_and_clock(event, cues)
    version = event.get("version", 1)

    wb = Workbook()
    ws = wb.active
    ws.title = "큐시트"

    ncol = len(COLUMNS)
    last_col = get_column_letter(ncol)

    # 컬럼 너비
    for idx, (_label, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    # Row 1: 타이틀
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = event.get("title", "(외부 주입 - 행사명)")
    c.fill, c.font, c.alignment = FILL_TITLE, FONT_TITLE, AL_C
    ws.row_dimensions[1].height = 30

    # Row 2: 메타
    ws.merge_cells(f"A2:{last_col}2")
    gen = datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = (f"일자: {event.get('date','-')}    베뉴: {event.get('venue','-')}    "
            f"버전: v{version}    생성: {gen}    "
            f"시간: {event.get('start_time','-')}~{event.get('end_time','-')}")
    c = ws["A2"]
    c.value, c.fill, c.font, c.alignment = meta, FILL_TITLE, FONT_META, AL_L
    ws.row_dimensions[2].height = 20

    # Row 3: 컬럼 헤더
    for idx, (label, _w) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=idx, value=label)
        cell.fill, cell.font, cell.alignment, cell.border = FILL_HDR, FONT_HDR, AL_C, BORDER
    ws.row_dimensions[3].height = 22

    # Row 4+: cue 행
    row = 4
    for i, cue in enumerate(cues):
        is_break = any(b in str(cue.get("segment", "")) for b in _BREAK_SEGMENTS)
        ws.cell(row=row, column=1, value=f"C{i+1:02d}")
        ws.cell(row=row, column=2, value=clocks[i])
        for j, key in enumerate(CUE_KEYS, start=3):
            ws.cell(row=row, column=j, value=cue.get(key, "") or ("-" if key == "stage" else ""))
        # 스타일
        for idx in range(1, ncol + 1):
            cell = ws.cell(row=row, column=idx)
            cell.border = BORDER
            cell.alignment = AL_C if idx in (1, 2, 9) else AL_L
            if idx == 8 and cue.get("cue"):       # 연출 cue 강조
                cell.font = FONT_CUE
            elif idx == 9:                         # Owner bold
                cell.font = FONT_BODY_BOLD
            else:
                cell.font = FONT_BODY
            if is_break:
                cell.fill = FILL_BREAK
        ws.row_dimensions[row].height = 26
        row += 1

    ws.freeze_panes = "A4"  # 헤더 고정

    # ── 변경이력 시트 ──
    ws2 = wb.create_sheet("변경이력")
    for idx, (label, width) in enumerate([("버전", 10), ("일자", 14), ("변경", 60)], start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = width
        cell = ws2.cell(row=1, column=idx, value=label)
        cell.fill, cell.font, cell.alignment, cell.border = FILL_HDR, FONT_HDR, AL_C, BORDER
    if not changelog:
        changelog = [(f"v{version}", date.today().isoformat(), "초안 생성")]
    for r, entry in enumerate(changelog, start=2):
        ver, day, note = (list(entry) + ["", "", ""])[:3]
        for idx, val in enumerate((ver, day, note), start=1):
            cell = ws2.cell(row=r, column=idx, value=val)
            cell.font, cell.border = FONT_BODY, BORDER
            cell.alignment = AL_L if idx == 3 else AL_C

    wb.save(output_path)
    if verify:
        _verify(output_path, len(cues), len(changelog))
    return output_path


# ============================================================
# 무결성 검증 (재오픈)
# ============================================================
def _verify(output_path: str, n_cues: int, n_changes: int):
    wb = load_workbook(output_path)
    if "큐시트" not in wb.sheetnames or "변경이력" not in wb.sheetnames:
        raise AssertionError(f"시트 누락: {wb.sheetnames}")
    ws = wb["큐시트"]
    if ws["A3"].value != "Cue#":
        raise AssertionError(f"컬럼 헤더 손상: A3={ws['A3'].value!r}")
    # cue 행 수 = 전체 행 - 헤더 3행
    body = ws.max_row - 3
    if body != n_cues:
        raise AssertionError(f"cue 행 수 불일치: 기대 {n_cues}, 실제 {body}")
    wb.close()
    return True


# ============================================================
# 출력 ChainPayload (선택)
# ============================================================
def to_chain_payload(event: dict, cues: list, target: str = "mice-dashboard") -> dict:
    """큐시트 계획 → ChainPayload/v1 (source=mice-run-of-show)."""
    clocks = _validate_and_clock(event, cues)
    start = _to_min(event["start_time"])
    total = sum(c["duration_min"] for c in cues)
    return {
        "$schema": "ChainPayload/v1",
        "source": "mice-run-of-show",
        "version": "v1.0.0",
        "generatedAt": datetime.now().isoformat(),
        "target": target,
        "projectTitle": event.get("title", ""),
        "eventDate": event.get("date", ""),
        "plan": {
            "startTime": event.get("start_time"),
            "endTime": _to_clock(start + total),
            "totalMinutes": total,
            "cueCount": len(cues),
            "cues": [
                {"cueNo": f"C{i+1:02d}", "clock": clocks[i],
                 "segment": c["segment"], "durationMin": c["duration_min"],
                 "owner": c.get("owner", "")}
                for i, c in enumerate(cues)
            ],
        },
        "version_no": event.get("version", 1),
    }


# ============================================================
# CLI 자가검증 — 반일 컨퍼런스 (worked-example.md 와 동일 데이터)
# ============================================================
def _self_test() -> int:
    event = {
        "title": "T社 테크 포럼 2026", "date": "2026-06-20",
        "venue": "[그랜드볼룸]", "start_time": "09:00", "end_time": "12:30",
        "version": 1, "client_id": None,
    }
    cues = [
        {"segment": "등록·입장", "duration_min": 30, "stage": "-", "audio": "BGM", "video": "로비 루프", "light": "하우스", "cue": "LX1 하우스 100%", "owner": "운영", "note": "정시 개문"},
        {"segment": "개회 선언", "duration_min": 5, "stage": "사회자", "audio": "MIC1", "video": "타이틀", "light": "무대 FULL", "cue": "BGM FADE → MIC1", "owner": "사회", "note": ""},
        {"segment": "환영사", "duration_min": 10, "stage": "주최 대표", "audio": "MIC2", "video": "발표 PPT", "light": "무대 FULL", "cue": "SB 환영사", "owner": "무대", "note": ""},
        {"segment": "기조연설", "duration_min": 35, "stage": "기조연사", "audio": "MIC2", "video": "LED→발표 PPT", "light": "무대 FULL+스팟", "cue": "멘트 끝 → VT1 + 암전", "owner": "무대", "note": "T-30초 SB"},
        {"segment": "기조 Q&A", "duration_min": 10, "stage": "기조연사+사회", "audio": "MIC1+MIC2", "video": "중계", "light": "무대 FULL", "cue": "스팟 OFF", "owner": "사회", "note": ""},
        {"segment": "휴식", "duration_min": 15, "stage": "-", "audio": "BGM", "video": "로비 루프", "light": "하우스", "cue": "LX1 하우스", "owner": "운영", "note": "세션장 셋업"},
        {"segment": "세션 1", "duration_min": 30, "stage": "연사 A", "audio": "MIC2", "video": "발표 PPT", "light": "무대 FULL", "cue": "SB 연사 A", "owner": "무대", "note": ""},
        {"segment": "세션 2", "duration_min": 30, "stage": "연사 B", "audio": "MIC2", "video": "발표 PPT", "light": "무대 FULL", "cue": "SB 연사 B", "owner": "무대", "note": ""},
        {"segment": "패널 토론", "duration_min": 30, "stage": "패널 4인+모더레이터", "audio": "MIC1~5", "video": "중계 PIP", "light": "무대 FULL", "cue": "테이블 세팅 전환", "owner": "무대", "note": "의자 5"},
        {"segment": "시상·기념촬영", "duration_min": 10, "stage": "시상자+수상자", "audio": "MIC1", "video": "시상 VT2", "light": "무대 FULL", "cue": "VT2 + 단체 스팟", "owner": "운영", "note": "포토월"},
        {"segment": "폐회", "duration_min": 5, "stage": "사회자", "audio": "MIC1", "video": "클로징", "light": "하우스 전환", "cue": "BGM IN", "owner": "사회", "note": ""},
    ]
    total = sum(c["duration_min"] for c in cues)
    assert total == 210, f"self-test 데이터 오류: 소요 합 {total} ≠ 210"

    out = Path(__file__).parent / "_self_test_runsheet.xlsx"
    build_runsheet(event, cues, str(out), verify=True)

    # 출력 ChainPayload 스모크
    cp = to_chain_payload(event, cues)
    assert cp["plan"]["totalMinutes"] == 210 and cp["plan"]["cueCount"] == 11
    assert cp["plan"]["endTime"] == "12:30", cp["plan"]["endTime"]

    print(f"PASS: {out.name}  (cue {len(cues)}개, 소요 합 {total}분 = 09:00~12:30, 무결성 검증)")

    # 시간 무결성 실패 케이스도 검증 (불일치는 ValueError 여야 함)
    try:
        bad = dict(event); bad["end_time"] = "12:00"
        build_runsheet(bad, cues, str(out), verify=False)
        print("FAIL: 시간 불일치인데 통과함")
        return 1
    except ValueError:
        print("PASS: 시간 무결성 실패 케이스 정상 차단(ValueError)")

    out.unlink(missing_ok=True)  # 자가검증 산출물 즉시 삭제
    return 0


if __name__ == "__main__":
    sys.exit(_self_test())
