#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_script.py — 발표 대본 .docx 빌드 (pt-script v2.0)

python-docx 를 사용하여 notes-extraction.json + (선택) proposal-meta.json
을 입력으로 받아 jc-design-system 토큰이 적용된 발표 대본 .docx 를 생성한다.

Usage:
    python build_script.py \\
        --notes /tmp/notes.json \\
        --meta /tmp/proposal-meta.json \\
        --out /tmp/script.docx

    또는 (메타 JSON 없이 CLI 옵션 직접):

    python build_script.py \\
        --notes /tmp/notes.json \\
        --minutes 20 \\
        --presentation-type bidding_pt \\
        --tone formal \\
        --presenter-role "발표자" \\
        --audience-type "발주처 심사위원" \\
        --client-name "[발주처명]" \\
        --rfp-title "[행사명] 운영 제안" \\
        --out /tmp/script.docx

본 스크립트는 jc-design-mapping.md §2~7 의 토큰 매핑을 자동 적용한다.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

# UTF-8 stdout 표준
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


# ============================================================
# JC Design System 토큰 (jc-design-mapping.md §2-3 기준)
# ============================================================
# docx는 CSS 변수를 못 쓰므로 HEX 리터럴은 "매체 불가피"로 남되,
# 값은 jc-design-system signature-tokens.md §6 JSON 라이트 정본과 일치시킨다 (SoT 미러).

# 컬러 토큰 (HEX) — jc-design-system(SoT) §6 JSON 런타임 로딩, 실패 시 미러 폴백.
def _load_jc_tokens():
    import json, re
    from pathlib import Path
    try:
        sot = Path(__file__).resolve().parents[2] / "jc-design-system" / "references" / "signature-tokens.md"
        m = re.search(r"```json\s*\n(.*?)\n```", sot.read_text(encoding="utf-8"), re.S)
        return json.loads(m.group(1)) if m else {}
    except Exception:
        return {}

_JC = _load_jc_tokens()
def _jc(key, fallback):
    c = _JC.get("color", {})
    v = c.get(key) or c.get("point", {}).get(key) or c.get("semantic", {}).get(key) or fallback
    return v.lstrip("#")

JC_PRIMARY = _jc("primary", "0A2540")              # --jc-primary
JC_PRIMARY_SOFT = _jc("primarySoft", "1A3556")     # --jc-primary-soft
JC_ACCENT = _jc("accent", "2962FF")                # --jc-accent
JC_ACCENT_STRONG = _jc("accentStrong", "1E4DCC")   # --jc-accent-strong
JC_TEXT = _jc("text", "1A1D24")                    # --jc-text
JC_TEXT_MUTED = _jc("textMuted", "5A6270")         # --jc-text-muted
JC_SURFACE = _jc("surface", "FFFFFF")              # --jc-surface
JC_SURFACE_ALT = _jc("surfaceAlt", "F1F3F7")       # --jc-surface-alt
JC_BORDER = _jc("border", "E5E8ED")                # --jc-border
JC_BORDER_STRONG = _jc("borderStrong", "C9CFD8")   # --jc-border-strong
JC_SUCCESS = _jc("success", "00C853")              # --jc-success
JC_WARNING = _jc("warning", "FFA000")              # --jc-warning
JC_DANGER = _jc("danger", "D32F2F")                # --jc-danger
JC_POINT_ORANGE = _jc("orange", "FF5722")          # --jc-point-orange
# mc 표지용 옅은 오렌지 배경 (--jc-point-orange-softest). 구 drift값 FFE5DD → FFF3E0 정합 (대비 14.8:1 AAA).
JC_POINT_ORANGE_SOFTEST = _jc("orangeSoftest", "FFF3E0")  # --jc-point-orange-softest

# 발표 유형별 강조 컬러 매핑 (jc-design-mapping.md §5)
PRESENTATION_TYPE_OVERRIDES = {
    "bidding_pt": {"heading_accent": JC_ACCENT, "cover_bg": JC_PRIMARY},
    "conference": {"heading_accent": JC_PRIMARY_SOFT, "cover_bg": JC_PRIMARY},
    "forum": {"heading_accent": JC_PRIMARY_SOFT, "cover_bg": JC_PRIMARY},
    "corporate_event": {"heading_accent": JC_POINT_ORANGE, "cover_bg": JC_PRIMARY},
    "mc": {"heading_accent": JC_POINT_ORANGE, "cover_bg": JC_POINT_ORANGE_SOFTEST},
    "general_business": {"heading_accent": JC_ACCENT, "cover_bg": JC_PRIMARY},
}

# 회사 종속 표현 (검증용)
FORBIDDEN_COMPANY_TERMS = [
    "엠앤씨", "M&C커뮤니케이션즈", "리멤버앤컴퍼니", "신사업실",
]


# ============================================================
# python-docx 헬퍼 — OXML 직접 조작
# ============================================================

def hex_to_rgb(hex_str: str):
    """HEX 문자열 → RGBColor 객체."""
    from docx.shared import RGBColor
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, font_name: str = "Pretendard", size_pt: Optional[int] = None,
                 color_hex: Optional[str] = None, bold: bool = False,
                 italic: bool = False):
    """Run 폰트·사이즈·컬러·스타일 일괄 설정 (CJK EastAsia 포함)."""
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    run.font.name = font_name
    # EastAsia 폰트 명시 (한글 폴백 보장)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)

    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color_hex is not None:
        run.font.color.rgb = hex_to_rgb(color_hex)
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def set_cell_shading(cell, hex_color: str):
    """셀 음영 적용 (jc-design-mapping.md §2-4)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_paragraph_shading(paragraph, hex_color: str):
    """단락 배경 음영 적용 (표지 헤더 등)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    p_pr.append(shd)


def add_page_break(doc):
    from docx.enum.text import WD_BREAK
    run = doc.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)


# ============================================================
# 시간 배분 로직 (SKILL.md Phase 3)
# ============================================================

SLIDE_TYPE_WEIGHTS = {
    "cover": 0.3,
    "agenda": 0.3,
    "thanks": 0.3,
    "section_divider": 0.3,
    "core_proposal": 1.5,
    "data_chart": 1.5,
    "background": 1.0,
    "process_timeline": 1.0,
    "reference_case": 1.0,
    "generic": 1.0,
    "image_centric": 0.7,
}

CHARS_PER_MINUTE = 250  # 한국어 기준


def calculate_time_allocation(slides: list, total_minutes: int, qna_minutes: int = 0) -> list:
    """슬라이드별 배정 시간 계산.

    Returns:
        list of dict — [{index, title, slide_type, weight, seconds, cumulative_seconds}]
    """
    presentation_seconds = (total_minutes - qna_minutes) * 60
    if presentation_seconds <= 0:
        presentation_seconds = total_minutes * 60  # 안전망

    # 가중치 합계
    total_weight = sum(
        SLIDE_TYPE_WEIGHTS.get(s["slide_type_hint"], 1.0) for s in slides
    )
    if total_weight == 0:
        total_weight = 1

    allocations = []
    cumulative = 0
    for s in slides:
        weight = SLIDE_TYPE_WEIGHTS.get(s["slide_type_hint"], 1.0)
        seconds = int(round(presentation_seconds * weight / total_weight))
        cumulative += seconds
        allocations.append({
            "index": s["index"],
            "title": s["title"] or f"슬라이드 {s['index']}",
            "slide_type": s["slide_type_hint"],
            "weight": weight,
            "seconds": seconds,
            "cumulative_seconds": cumulative,
        })

    return allocations


def format_seconds(sec: int) -> str:
    """초 → 'X분 Y초' 형식."""
    m, s = divmod(sec, 60)
    if m == 0:
        return f"{s}초"
    if s == 0:
        return f"{m}분"
    return f"{m}분 {s}초"


# ============================================================
# 회사 종속 표현 검증
# ============================================================

def validate_company_mentions(text: str) -> list:
    """텍스트에서 회사 종속 표현 검출."""
    detected = []
    for term in FORBIDDEN_COMPANY_TERMS:
        if term in text:
            detected.append(term)
    return detected


def sanitize_text(text: str) -> str:
    """회사 종속 표현을 일반 표현으로 치환."""
    if not text:
        return ""
    sanitized = text
    sanitized = re.sub(r"엠앤씨\s*커뮤니케이션즈", "[발표 주체]", sanitized)
    sanitized = re.sub(r"M&C\s*커뮤니케이션즈", "[발표 주체]", sanitized)
    sanitized = re.sub(r"리멤버앤컴퍼니", "[발표 주체]", sanitized)
    sanitized = re.sub(r"신사업실", "[부서명]", sanitized)
    return sanitized


# ============================================================
# 멘트 생성 (간이 자동 생성 — 정식 LLM 호출은 상위 워크플로우에서)
# ============================================================

def generate_mention(slide: dict, allocation: dict, tone: str = "formal",
                     presentation_type: str = "bidding_pt") -> str:
    """슬라이드 데이터 기반 발표 멘트 자동 생성 (폴백용 간이 버전).

    실제 운영 시 LLM 호출로 더 자연스러운 멘트 생성 권장.
    본 함수는 스크립트 단독 실행 시의 폴백.
    """
    notes_type = slide.get("notes_type", "empty")
    title = slide.get("title", "")
    body_text = slide.get("body_text", "")
    notes = slide.get("notes", "")
    slide_type = slide.get("slide_type_hint", "generic")

    # 우선순위 1: speaker notes 그대로 활용 (sanitize 만 적용)
    if notes_type == "speaker" and notes:
        return sanitize_text(notes)

    # 우선순위 2~3: 슬라이드 유형별 패턴
    if slide_type == "cover":
        if presentation_type == "bidding_pt":
            return f"안녕하십니까. 오늘 {title} 에 대한 저희의 제안을 말씀드리겠습니다. " \
                   f"약 {allocation['seconds']//60}분간 발표 후 질의응답 시간을 갖겠습니다."
        if presentation_type == "mc":
            return f"안녕하십니까, {title} 에 참석해 주신 여러분을 진심으로 환영합니다."
        return f"안녕하십니까. 오늘 {title} 주제로 말씀드리겠습니다."

    if slide_type == "agenda":
        items = [l.strip() for l in body_text.split("\n") if l.strip()]
        if items:
            return f"오늘 발표는 크게 {len(items)} 가지로 구성했습니다. " + \
                   ", ".join(items[:3]) + " 등을 순서대로 말씀드리겠습니다."
        return "오늘 발표 순서를 간략히 안내드리겠습니다."

    if slide_type == "thanks":
        if presentation_type == "bidding_pt":
            return "이상으로 제안 발표를 마치겠습니다. 경청해 주셔서 감사합니다. 질의응답을 받겠습니다."
        return "경청해 주셔서 감사합니다."

    if slide_type == "data_chart":
        return f"화면의 데이터를 보시면, {title} 와 관련된 핵심 수치를 확인할 수 있습니다. " \
               f"{body_text[:100]}..." if body_text else f"{title} 의 데이터를 보시겠습니다."

    if slide_type == "core_proposal":
        return f"저희가 제안드리는 핵심은 {title} 입니다. {body_text[:150]}" if body_text \
               else f"저희가 제안드리는 핵심 전략 {title} 에 대해 말씀드리겠습니다."

    if slide_type == "reference_case":
        return f"저희의 관련 실적을 말씀드리겠습니다. {body_text[:150]}" if body_text \
               else f"{title} 와 관련된 저희 실적을 소개드리겠습니다."

    if slide_type == "background":
        return f"먼저 {title} 의 현재 상황을 짚어보겠습니다. {body_text[:150]}" if body_text \
               else f"{title} 에 대한 배경을 말씀드리겠습니다."

    if slide_type == "process_timeline":
        return f"전체 {title} 을 단계별로 정리했습니다. {body_text[:150]}" if body_text \
               else f"{title} 의 진행 단계를 설명드리겠습니다."

    # 폴백
    if body_text:
        return f"{title} 에 대해 말씀드리겠습니다. {body_text[:150]}"
    return f"{title} 에 대해 말씀드리겠습니다."


def generate_tip(slide_type: str) -> str:
    """슬라이드 유형별 발표 팁 (script-guide.md 의 표 참조)."""
    tips = {
        "cover": "당당한 자세, 청중과 눈 맞춤. 원고를 보지 말 것.",
        "data_chart": "레이저 포인터 또는 손으로 핵심 수치를 가리키며 설명.",
        "core_proposal": "속도를 약간 늦추고, 목소리 톤을 높여 강조.",
        "reference_case": "에피소드나 스토리를 곁들여 생동감 부여.",
        "thanks": "마지막 핵심 메시지를 천천히, 명확하게 전달.",
        "agenda": "각 항목을 또렷이 말하고, 손가락으로 순서 표시.",
    }
    return tips.get(slide_type, "슬라이드 전환 시 1~2초 멈춤(pause) 으로 청중 집중 유도.")


# ============================================================
# Q&A 자동 생성 (간이)
# ============================================================

def generate_qna(slides: list, presentation_type: str = "bidding_pt") -> list:
    """슬라이드 기반 예상 Q&A 자동 생성 (간이 버전)."""
    qna: list[dict] = []

    if presentation_type == "bidding_pt":
        qna.extend([
            {
                "q": "제안 내용을 실제로 약속하신 기간 안에 구현할 수 있는 근거는 무엇입니까?",
                "a": "유사 규모의 행사를 성공적으로 운영한 다수의 실적이 있으며, "
                     "본 제안서에서 명시한 일정은 해당 노하우 기반 산출 일정입니다."
            },
            {
                "q": "타 업체와 비교했을 때 귀사만의 가장 큰 강점은 무엇입니까?",
                "a": "본 제안서 핵심 전략 슬라이드에서 강조한 3가지 차별점 — "
                     "운영 안정성, 디지털 확장성, 사후 관리 — 가 핵심 강점입니다."
            },
            {
                "q": "예산 산출 근거를 항목별로 설명해 주시기 바랍니다.",
                "a": "각 항목은 국가계약법 산출 기준과 시장 표준 단가를 적용했으며, "
                     "별도 견적 부록에 항목별 단가 출처를 명시했습니다."
            },
            {
                "q": "핵심 인력의 전담 여부와 경력을 확인하고 싶습니다.",
                "a": "본 행사에는 PM 1명, 운영 매니저 2명이 전담 투입되며, "
                     "모두 유사 규모 행사 운영 경력 5년 이상입니다."
            },
            {
                "q": "예상치 못한 리스크 발생 시 대응 방안은 어떻게 됩니까?",
                "a": "본 제안서 리스크 관리 섹션에 명시한 3단계 대응 프로토콜 — "
                     "사전 예방·실시간 모니터링·사후 복구 — 을 통해 대응합니다."
            },
        ])
    else:
        qna.extend([
            {
                "q": "발표 내용의 구체적인 방법론은 어떻게 됩니까?",
                "a": "각 슬라이드에서 설명드린 단계를 순차적으로 적용한 방법론입니다."
            },
            {
                "q": "제시하신 데이터의 출처와 신뢰성에 대해 설명 부탁드립니다.",
                "a": "본 발표에 인용한 데이터는 공인 기관 통계와 자체 조사 결과를 결합한 것입니다."
            },
            {
                "q": "이 사례가 다른 분야에도 적용 가능합니까?",
                "a": "기본 원칙은 동일하게 적용 가능하며, 분야별 특성에 맞춰 일부 조정이 필요합니다."
            },
        ])

    return qna


# ============================================================
# docx 빌드 메인
# ============================================================

def build_script_docx(
    notes_data: dict,
    meta: dict,
    output_path: str,
) -> dict:
    """발표 대본 .docx 생성.

    Args:
        notes_data: extract_notes.py 출력 JSON
        meta: proposal_meta (필수 필드: presentation_minutes, presentation_type, tone, ...)
        output_path: 출력 .docx 경로

    Returns:
        dict — {output_path, slide_count, total_chars, estimated_minutes, time_overrun}
    """
    try:
        from docx import Document
        from docx.shared import Pt, Cm, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
    except ImportError as exc:
        raise RuntimeError(
            "python-docx 가 설치되지 않았습니다. `pip install python-docx` 실행."
        ) from exc

    # ----- 메타 추출 -----
    total_minutes = int(meta.get("presentation_minutes", 20))
    qna_minutes = int(meta.get("qna_minutes", round(total_minutes * 0.2)))
    presentation_type = meta.get("presentation_type", "bidding_pt")
    tone = meta.get("tone", "formal")
    client_name = sanitize_text(meta.get("client_name", "[발주처명]"))
    rfp_title = sanitize_text(meta.get("rfp_title", "[발표 주제]"))
    presenter_role = sanitize_text(meta.get("presenter_role", "발표자"))
    presenter_name = sanitize_text(meta.get("presenter_name") or "")
    audience_type = meta.get("audience_type", "심사위원")
    qna_included = meta.get("qna_included", True)

    overrides = PRESENTATION_TYPE_OVERRIDES.get(
        presentation_type, PRESENTATION_TYPE_OVERRIDES["bidding_pt"]
    )
    cover_bg = overrides["cover_bg"]
    heading_accent = overrides["heading_accent"]

    # 표지 텍스트 컬러 (배경이 밝으면 진한 색, 어두우면 흰색)
    cover_text_color = JC_SURFACE if cover_bg in (JC_PRIMARY,) else JC_TEXT

    # ----- 시간 배분 계산 -----
    slides = notes_data.get("slides", [])
    allocations = calculate_time_allocation(
        slides, total_minutes, qna_minutes if qna_included else 0
    )

    # ----- Document 생성 -----
    doc = Document()

    # 페이지 마진 설정 (A4 기본 + 적정 마진)
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ===== [표지] =====
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_shading(title_para, cover_bg)
    title_para.paragraph_format.space_before = Pt(60)
    title_para.paragraph_format.space_after = Pt(60)

    title_run = title_para.add_run(rfp_title)
    set_run_font(title_run, "Pretendard", size_pt=28, color_hex=cover_text_color, bold=True)

    # 표지 메타 정보
    meta_lines = [
        f"발표 대상: {audience_type}",
        f"발표자: {presenter_role}" + (f" {presenter_name}" if presenter_name else ""),
        f"발표 시간: {total_minutes}분" + (f" (Q&A {qna_minutes}분 포함)" if qna_included else ""),
        f"발표 일자: {datetime.now().strftime('%Y년 %m월 %d일')}",
        f"발주처: {client_name}",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, "Pretendard", size_pt=12, color_hex=JC_TEXT_MUTED)

    add_page_break(doc)

    # ===== [발표 개요 — 시간 배분표] =====
    overview_heading = doc.add_paragraph()
    r = overview_heading.add_run("발표 개요 — 시간 배분")
    set_run_font(r, "Pretendard", size_pt=24, color_hex=JC_PRIMARY, bold=True)
    overview_heading.paragraph_format.space_after = Pt(18)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # 헤더
    hdr = table.rows[0].cells
    headers = ["슬라이드", "제목", "배정 시간", "누적 시간"]
    for i, h in enumerate(headers):
        hdr[i].text = ""  # 기존 텍스트 비우기
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, "Pretendard", size_pt=11, color_hex=JC_SURFACE, bold=True)
        set_cell_shading(hdr[i], JC_PRIMARY)

    # 데이터 행
    for i, alloc in enumerate(allocations):
        row = table.add_row().cells
        bg = JC_SURFACE_ALT if (i + 1) % 2 == 0 else JC_SURFACE

        row[0].text = ""
        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(f"#{alloc['index']}")
        set_run_font(r0, "Pretendard", size_pt=11, color_hex=JC_TEXT)
        if bg != JC_SURFACE:
            set_cell_shading(row[0], bg)

        row[1].text = ""
        p1 = row[1].paragraphs[0]
        r1 = p1.add_run(alloc["title"][:30])
        set_run_font(r1, "Pretendard", size_pt=11, color_hex=JC_TEXT)
        if bg != JC_SURFACE:
            set_cell_shading(row[1], bg)

        row[2].text = ""
        p2 = row[2].paragraphs[0]
        r2 = p2.add_run(format_seconds(alloc["seconds"]))
        set_run_font(r2, "Pretendard", size_pt=11, color_hex=JC_TEXT)
        if bg != JC_SURFACE:
            set_cell_shading(row[2], bg)

        row[3].text = ""
        p3 = row[3].paragraphs[0]
        r3 = p3.add_run(format_seconds(alloc["cumulative_seconds"]))
        set_run_font(r3, "Pretendard", size_pt=11, color_hex=JC_TEXT)
        if bg != JC_SURFACE:
            set_cell_shading(row[3], bg)

    add_page_break(doc)

    # ===== [슬라이드별 스크립트] =====
    section_heading = doc.add_paragraph()
    r = section_heading.add_run("슬라이드별 발표 스크립트")
    set_run_font(r, "Pretendard", size_pt=24, color_hex=JC_PRIMARY, bold=True)
    section_heading.paragraph_format.space_after = Pt(18)

    total_chars = 0
    for slide, alloc in zip(slides, allocations):
        # 슬라이드 헤딩 (Heading2 강조)
        sh = doc.add_paragraph()
        r = sh.add_run(f"슬라이드 {slide['index']}: {slide.get('title') or '(제목 없음)'}")
        set_run_font(r, "Pretendard", size_pt=20, color_hex=heading_accent, bold=True)
        sh.paragraph_format.space_before = Pt(18)
        sh.paragraph_format.space_after = Pt(6)

        # 배정 시간
        time_p = doc.add_paragraph()
        r = time_p.add_run(f"배정 시간: {format_seconds(alloc['seconds'])}  |  유형: {alloc['slide_type']}")
        set_run_font(r, "Pretendard", size_pt=11, color_hex=JC_TEXT_MUTED, italic=True)
        time_p.paragraph_format.space_after = Pt(6)

        # 슬라이드 내용 요약
        if slide.get("body_text"):
            summary_label = doc.add_paragraph()
            r = summary_label.add_run("[슬라이드 내용 요약]")
            set_run_font(r, "Pretendard", size_pt=12, color_hex=JC_TEXT, bold=True)
            summary_label.paragraph_format.space_after = Pt(2)

            summary_p = doc.add_paragraph()
            r = summary_p.add_run(sanitize_text(slide["body_text"])[:300])
            set_run_font(r, "Pretendard", size_pt=11, color_hex=JC_TEXT_MUTED)
            summary_p.paragraph_format.line_spacing = 1.3
            summary_p.paragraph_format.space_after = Pt(10)

        # 발표 멘트 (핵심 영역)
        mention_label = doc.add_paragraph()
        r = mention_label.add_run("【발표 멘트】")
        set_run_font(r, "Pretendard", size_pt=14, color_hex=JC_ACCENT_STRONG, bold=True)
        mention_label.paragraph_format.space_after = Pt(4)

        mention_text = generate_mention(slide, alloc, tone=tone, presentation_type=presentation_type)
        mention_text = sanitize_text(mention_text)
        total_chars += len(mention_text)

        mention_p = doc.add_paragraph()
        r = mention_p.add_run(mention_text)
        set_run_font(r, "Pretendard", size_pt=14, color_hex=JC_TEXT)
        mention_p.paragraph_format.line_spacing = 1.5
        mention_p.paragraph_format.space_after = Pt(10)

        # 발표 팁
        tip_p = doc.add_paragraph()
        r = tip_p.add_run(f"💡 발표 팁: {generate_tip(alloc['slide_type'])}")
        set_run_font(r, "Pretendard", size_pt=11, color_hex=JC_TEXT_MUTED, italic=True)
        tip_p.paragraph_format.line_spacing = 1.3
        tip_p.paragraph_format.space_after = Pt(18)

    add_page_break(doc)

    # ===== [Q&A 예상 질의응답] =====
    if qna_included:
        qna_heading = doc.add_paragraph()
        r = qna_heading.add_run("Q&A 예상 질의응답")
        set_run_font(r, "Pretendard", size_pt=24, color_hex=JC_PRIMARY, bold=True)
        qna_heading.paragraph_format.space_after = Pt(18)

        qna_list = generate_qna(slides, presentation_type)
        for i, item in enumerate(qna_list, start=1):
            # 질문
            q_p = doc.add_paragraph()
            r = q_p.add_run(f"Q{i}. {sanitize_text(item['q'])}")
            set_run_font(r, "Pretendard", size_pt=14, color_hex=JC_TEXT, bold=True)
            q_p.paragraph_format.space_before = Pt(12)
            q_p.paragraph_format.space_after = Pt(4)

            # 답변 라벨
            a_label = doc.add_paragraph()
            r = a_label.add_run("예상 답변:")
            set_run_font(r, "Pretendard", size_pt=12, color_hex=JC_SUCCESS, bold=True)
            a_label.paragraph_format.space_after = Pt(2)

            # 답변 본문
            a_p = doc.add_paragraph()
            r = a_p.add_run(sanitize_text(item["a"]))
            set_run_font(r, "Pretendard", size_pt=14, color_hex=JC_TEXT)
            a_p.paragraph_format.line_spacing = 1.5
            a_p.paragraph_format.space_after = Pt(12)

        add_page_break(doc)

    # ===== [발표 체크리스트] =====
    cl_heading = doc.add_paragraph()
    r = cl_heading.add_run("발표 체크리스트")
    set_run_font(r, "Pretendard", size_pt=24, color_hex=JC_PRIMARY, bold=True)
    cl_heading.paragraph_format.space_after = Pt(18)

    checklist = {
        "발표 전 (24시간 이내)": [
            "원고 최종 리허설 1회 이상 완료",
            "발표 시간 ±10% 이내 확인",
            "Q&A 답변 키 메시지 숙지",
            "장비 점검 (마이크·노트북·포인터)",
            "발표 자료 백업본 USB 준비",
        ],
        "발표 중 (현장)": [
            "첫 30초 청중과 눈 맞춤",
            "슬라이드 전환 시 1~2초 호흡",
            "핵심 수치 손가락·포인터로 가리키며 강조",
            "시간 배분표 vs 실제 진행 시간 모니터링",
            "Q&A 질문은 끝까지 듣고 답변",
        ],
        "발표 후": [
            "심사위원·청중 명함 교환",
            "발표 자료·Q&A 답변 요약 메일 전달",
            "내부 디브리프 — 잘된 점·개선점 기록",
        ],
    }
    for category, items in checklist.items():
        cat_p = doc.add_paragraph()
        r = cat_p.add_run(f"■ {category}")
        set_run_font(r, "Pretendard", size_pt=14, color_hex=heading_accent, bold=True)
        cat_p.paragraph_format.space_before = Pt(12)
        cat_p.paragraph_format.space_after = Pt(4)

        for item in items:
            li = doc.add_paragraph()
            r = li.add_run(f"□  {item}")
            set_run_font(r, "Pretendard", size_pt=12, color_hex=JC_TEXT)
            li.paragraph_format.left_indent = Inches(0.3)
            li.paragraph_format.line_spacing = 1.4

    # ===== 글자 수 → 예상 시간 검증 =====
    estimated_minutes = round(total_chars / CHARS_PER_MINUTE, 1)
    pure_minutes = total_minutes - (qna_minutes if qna_included else 0)
    overrun_pct = abs(estimated_minutes - pure_minutes) / max(pure_minutes, 1) * 100
    time_overrun = overrun_pct > 10

    # 시간 초과 시 마지막에 경고 단락 추가
    if time_overrun:
        warning_p = doc.add_paragraph()
        warning_p.paragraph_format.space_before = Pt(24)
        set_paragraph_shading(warning_p, JC_WARNING)
        r = warning_p.add_run(
            f"⚠ 시간 검증: 예상 멘트 시간 {estimated_minutes}분 vs 순수 발표 시간 {pure_minutes}분. "
            f"오차 {overrun_pct:.1f}% — 리허설 시 멘트 조정 권장."
        )
        set_run_font(r, "Pretendard", size_pt=12, color_hex=JC_TEXT, bold=True)

    # ===== 저장 =====
    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))

    return {
        "output_path": str(out_path),
        "slide_count": len(slides),
        "total_chars": total_chars,
        "estimated_minutes": estimated_minutes,
        "pure_minutes": pure_minutes,
        "overrun_pct": round(overrun_pct, 1),
        "time_overrun": time_overrun,
    }


# ============================================================
# CLI 진입점
# ============================================================

def build_meta_from_args(args) -> dict:
    """CLI 인자로 proposal_meta 객체 구성 (메타 JSON 없을 시)."""
    return {
        "client_name": args.client_name or "[발주처명]",
        "rfp_title": args.rfp_title or "[발표 주제]",
        "presentation_minutes": args.minutes,
        "presentation_type": args.presentation_type,
        "presenter_role": args.presenter_role or "발표자",
        "presenter_name": args.presenter_name or "",
        "audience_type": args.audience_type or "심사위원",
        "tone": args.tone,
        "qna_included": not args.no_qna,
        "qna_minutes": args.qna_minutes if args.qna_minutes is not None else round(args.minutes * 0.2),
        "client_id": args.client_id,
    }


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="발표 대본 .docx 빌드 (pt-script v2.0)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--notes", required=True, help="extract_notes.py 산출 JSON 경로")
    parser.add_argument("--meta", default=None, help="proposal-meta.json 경로 (선택)")
    parser.add_argument("--out", "-o", required=True, help="출력 .docx 파일 경로")

    # 메타 JSON 없을 시 직접 옵션
    parser.add_argument("--minutes", type=int, default=20, help="전체 발표 시간 (분)")
    parser.add_argument(
        "--presentation-type",
        default="bidding_pt",
        choices=["bidding_pt", "conference", "forum", "corporate_event", "mc", "general_business"],
    )
    parser.add_argument("--tone", default="formal", choices=["formal", "semi_formal", "casual"])
    parser.add_argument("--client-name", default=None)
    parser.add_argument("--rfp-title", default=None)
    parser.add_argument("--presenter-role", default=None)
    parser.add_argument("--presenter-name", default=None)
    parser.add_argument("--audience-type", default=None)
    parser.add_argument("--qna-minutes", type=int, default=None)
    parser.add_argument("--no-qna", action="store_true", help="Q&A 섹션 제외")
    parser.add_argument("--client-id", default=None, help="클라이언트 오버레이 ID")

    parser.add_argument("--verbose", "-v", action="store_true")

    args = parser.parse_args(argv)

    # ----- 입력 로드 -----
    notes_path = Path(args.notes)
    if not notes_path.exists():
        print(f"[ERROR] notes JSON 파일 없음: {args.notes}", file=sys.stderr)
        return 2
    notes_data = json.loads(notes_path.read_text(encoding="utf-8"))

    if args.meta:
        meta_path = Path(args.meta)
        if not meta_path.exists():
            print(f"[ERROR] meta JSON 파일 없음: {args.meta}", file=sys.stderr)
            return 2
        meta_full = json.loads(meta_path.read_text(encoding="utf-8"))
        meta = meta_full.get("proposal_meta", {})
        # 메타 JSON 의 필드와 CLI 옵션 병합 (CLI 우선)
        if args.minutes != 20:  # 기본값 아니면 덮어쓰기
            meta["presentation_minutes"] = args.minutes
    else:
        meta = build_meta_from_args(args)

    # 회사 종속 표현 사전 검증
    company_warns = []
    for k, v in meta.items():
        if isinstance(v, str):
            detected = validate_company_mentions(v)
            for term in detected:
                company_warns.append(f"meta.{k} 에 '{term}' 검출 (자동 sanitize 적용)")
    if company_warns and args.verbose:
        for w in company_warns:
            print(f"[WARN] {w}", file=sys.stderr)

    # ----- 빌드 -----
    try:
        result = build_script_docx(notes_data, meta, args.out)
    except RuntimeError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 5
    except Exception as exc:
        print(f"[ERROR] 빌드 실패: {exc}", file=sys.stderr)
        return 6

    # ----- 결과 출력 -----
    summary = {
        "output_path": result["output_path"],
        "slide_count": result["slide_count"],
        "total_chars": result["total_chars"],
        "estimated_minutes": result["estimated_minutes"],
        "pure_minutes": result["pure_minutes"],
        "overrun_pct": result["overrun_pct"],
        "time_overrun": result["time_overrun"],
        "company_mention_warnings": company_warns,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if args.verbose:
        print(f"\n[OK] 빌드 완료 — {result['output_path']}", file=sys.stderr)
        print(f"     슬라이드: {result['slide_count']}개", file=sys.stderr)
        print(f"     멘트 총량: {result['total_chars']:,}자 ≒ {result['estimated_minutes']}분", file=sys.stderr)
        if result["time_overrun"]:
            print(f"     ⚠ 시간 오차 {result['overrun_pct']}% — 조정 권장", file=sys.stderr)

    return 0


if __name__ == "__main__":
    sys.exit(main())
